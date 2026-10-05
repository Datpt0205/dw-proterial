"""The committed attachments are exactly what the generator makes from the fixtures."""

from __future__ import annotations

import io
import re
import uuid
from collections.abc import Callable
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from pypdf import PdfReader
from pypdf.generic import DictionaryObject

from dw_kernel.ids import TenantId, WorkspaceId
from dw_sales.adapters.mock import MockInbox
from dw_sales.adapters.mock.fixtures import ATTACHMENTS_DIR, attachment_file_name
from dw_sales.adapters.mock.generate_attachments import (
    OrderLine,
    PurchaseOrderDocument,
    StatedAttributes,
    load_design_replies,
    load_purchase_orders,
    load_quote_requests,
    render_attachments,
)
from dw_sales.application.ports import SalesScope

pytestmark = pytest.mark.unit

SCOPE = SalesScope(TenantId(uuid.UUID(int=1)), WorkspaceId(uuid.UUID(int=2)))
REGENERATE = "uv run python -m dw_sales.adapters.mock.generate_attachments"


@pytest.fixture(scope="module")
def rendered() -> dict[str, bytes]:
    return render_attachments()


def test_the_committed_attachments_are_what_the_generator_makes(rendered: dict[str, bytes]) -> None:
    committed = {path.name: path.read_bytes() for path in ATTACHMENTS_DIR.iterdir()}

    assert sorted(committed) == sorted(rendered), f"files differ from the fixtures: {REGENERATE}"
    stale = sorted(name for name, data in rendered.items() if committed[name] != data)
    assert not stale, f"{stale} are stale: {REGENERATE}"


def test_rendering_twice_gives_the_same_bytes(rendered: dict[str, bytes]) -> None:
    assert render_attachments() == rendered


async def test_every_attachment_in_the_inbox_has_a_document_and_no_document_is_orphaned(
    rendered: dict[str, bytes],
) -> None:
    inbox = MockInbox.load(SCOPE)
    carried = {
        attachment_file_name(message.message_id, attachment.name)
        for message in await inbox.list_messages(SCOPE)
        for attachment in message.attachments
    }

    assert carried == set(rendered)


def test_the_pdf_text_layer_carries_vietnamese_one_row_per_line() -> None:
    (po,) = [p for p in load_purchase_orders() if p.layout == "pdf"]
    reader = PdfReader(ATTACHMENTS_DIR / attachment_file_name(po.message_id, po.name))
    first, terms = (page.extract_text(extraction_mode="layout") for page in reader.pages)
    # Layout mode puts each printed row on one line, cells two or more spaces apart.
    rows = {
        cells[0]: cells[1:]
        for cells in (re.split(r"\s{2,}", line.strip()) for line in first.splitlines())
    }

    assert "ĐƠN ĐẶT HÀNG (PURCHASE ORDER)" in first
    assert "ĐIỀU KHOẢN CHUNG" in terms
    assert [rows[str(line.no)][0] for line in po.lines] == [
        line.customer_item_code for line in po.lines
    ]
    assert rows["2"] == [
        "BRN-C-0112",
        "MULTI-CORE CABLE 2C 0.75mm2 STRANDED BARE CU NO SHIELD GREY REEL 300M SPEC SP-5202",
        "3.000",
        "m",
        "14.500",
        "43.500.000",
        "30/11/2026",
    ]


def test_the_intra_group_po_prints_one_sheet_per_page(rendered: dict[str, bytes]) -> None:
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M02"]
    workbook = load_workbook(io.BytesIO(rendered[attachment_file_name(po.message_id, po.name)]))

    assert workbook.sheetnames == ["Page 1", "Page 2", "Page 3", "Page 4"]
    numbers = []
    for page, sheet in enumerate(workbook.worksheets, start=1):
        assert sheet["B4"].value == po.po_no
        assert sheet["E6"].value == f"{page} / 4"
        numbers += [
            row[0].value for row in sheet.iter_rows(min_row=10) if isinstance(row[0].value, int)
        ]
        totals = [
            cell for row in sheet.iter_rows(min_row=10) for cell in row if cell.value == "TOTAL"
        ]
        assert len(totals) == (1 if page == 4 else 0)
    assert numbers == list(range(1, 19))


def test_a_mapped_line_is_described_as_the_catalogue_describes_its_item(
    rendered: dict[str, bytes],
) -> None:
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M01"]
    workbook = load_workbook(io.BytesIO(rendered[attachment_file_name(po.message_id, po.name)]))

    assert workbook["PO"]["C10"].value == (
        "HOOK-UP WIRE 1C AWG24 STRANDED TINNED CU NO SHIELD BLACK REEL 610M SPEC SP-1124"
    )


@pytest.mark.parametrize(
    ("change", "message"),
    [
        (lambda po: po | {"lines_per_page": 5}, "sheet-per-page layout only"),
        (lambda po: po | {"layout": "xlsx_sheet_per_page"}, "sheet-per-page layout only"),
        (lambda po: po | {"terms": ["Payment in 60 days."]}, "printed by the PDF only"),
        (lambda po: po | {"buyer_address": "Lot 1, Demo Park"}, "printed by the PDF only"),
        (lambda po: po | {"lines": po["lines"][::-1]}, "numbered in increasing order"),
        (lambda po: po | {"buyer_name": "Someone Else Ltd."}, "one of them"),
        (lambda po: po | {"customer_code": None}, "one of them"),
        (lambda po: po | {"hidden_lines": [9]}, "a line the PO lacks"),
        (lambda po: po | {"hidden_pages": [1]}, "no page before the last"),
        (
            lambda po: po | {"layout": "pdf", "hidden_lines": [1]},
            "only an Excel layout hides",
        ),
    ],
    ids=[
        "pages_on_one_sheet",
        "pages_unsized",
        "terms_in_excel",
        "address_in_excel",
        "order",
        "customer_and_buyer",
        "neither",
        "hidden_line_missing",
        "hiding_the_total",
        "hidden_in_a_pdf",
    ],
)
def test_a_document_its_layout_cannot_print_is_refused(
    change: Callable[[dict[str, Any]], dict[str, Any]], message: str
) -> None:
    """Otherwise a fixture edit changes nothing in the file, or a README line
    number points at another line."""
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M01"]

    with pytest.raises(ValidationError, match=message):
        PurchaseOrderDocument.model_validate(change(po.model_dump()))


def test_unstated_attributes_are_left_out_of_a_description() -> None:
    stated = StatedAttributes(
        family="multi_core_cable",
        cores=2,
        gauge="0.5mm2",
        stranding="stranded",
        conductor="tinned_copper",
        shield="foil",
        colour="black",
        spec_no="SP-5201",
    )

    assert (
        stated.describe()
        == "MULTI-CORE CABLE 2C 0.5mm2 STRANDED TINNED CU FOIL SHIELD BLACK SPEC SP-5201"
    )


def _workbook(rendered: dict[str, bytes], message_id: str) -> Workbook:
    names = (
        {d.message_id: d.name for d in load_purchase_orders()}
        | {d.message_id: d.name for d in load_quote_requests()}
        | {d.message_id: d.name for d in load_design_replies()}
    )
    return load_workbook(io.BytesIO(rendered[attachment_file_name(message_id, names[message_id])]))


def test_the_scanned_po_is_one_image_and_no_text(rendered: dict[str, bytes]) -> None:
    """What makes it unreadable is the absence of a text layer, so that is
    what is held: no text in either extraction mode, no font, one image."""
    (po,) = [p for p in load_purchase_orders() if p.layout == "pdf_image"]
    reader = PdfReader(io.BytesIO(rendered[attachment_file_name(po.message_id, po.name)]))
    (page,) = reader.pages

    assert page.extract_text() == ""
    assert page.extract_text(extraction_mode="layout") == ""
    resources = page["/Resources"]
    assert isinstance(resources, DictionaryObject)
    assert "/Font" not in resources
    assert len(page.images) == 1


def test_text_that_starts_with_an_equals_sign_is_stored_as_text(
    rendered: dict[str, bytes],
) -> None:
    """Ticket 06's formula test needs the customer's text, not a formula without
    a cached value, which is another fixture (`value_uncertain`)."""
    cell = _workbook(rendered, "M20")["Page 1"]["C12"]

    assert cell.data_type == "s"
    assert isinstance(cell.value, str)
    assert cell.value.startswith("=")


def test_the_long_intra_group_po_spans_eleven_sheets_and_62_lines(
    rendered: dict[str, bytes],
) -> None:
    workbook = _workbook(rendered, "M20")
    numbers = [
        row[0].value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(min_row=10)
        if isinstance(row[0].value, int)
    ]

    assert len(workbook.worksheets) == 11
    assert numbers == list(range(1, 63))


def test_hidden_rows_and_sheets_are_hidden_in_the_file(rendered: dict[str, bytes]) -> None:
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M27"]
    workbook = _workbook(rendered, "M27")

    assert [s.sheet_state for s in workbook.worksheets] == ["visible", "hidden", "visible"]
    assert po.hidden_lines == (2,)
    assert workbook["Page 1"]["A11"].value == 2
    assert workbook["Page 1"].row_dimensions[11].hidden
    assert not workbook["Page 1"].row_dimensions[10].hidden


def test_a_printed_total_is_printed_as_the_fixture_states(rendered: dict[str, bytes]) -> None:
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M21"]
    sheet = _workbook(rendered, "M21")["PO"]

    assert po.printed_total is not None
    assert sheet["F12"].value == "TOTAL"
    assert Decimal(str(sheet["G12"].value)) == po.printed_total


def test_a_unit_is_printed_as_the_customer_wrote_it(rendered: dict[str, bytes]) -> None:
    assert _workbook(rendered, "M26")["PO"]["E11"].value == "FT"


def test_a_quantity_the_request_leaves_out_is_a_blank_cell(rendered: dict[str, bytes]) -> None:
    sheet = _workbook(rendered, "M23")["RFQ"]

    assert sheet["D9"].value == "Quantity"
    assert sheet["D10"].value is None


def test_a_design_reply_names_its_ycbg_and_what_design_specified(
    rendered: dict[str, bytes],
) -> None:
    (reply,) = [r for r in load_design_replies() if r.message_id == "M25"]
    sheet = _workbook(rendered, "M25")["YCBG"]

    assert sheet["B4"].value == reply.ycbg_no
    assert [cell.value for cell in sheet[10]][:4] == [1, "BP-25-0187", "SP-5406", "CB-2007"]


def test_a_buyer_that_is_no_customer_is_printed_by_its_own_name(
    rendered: dict[str, bytes],
) -> None:
    (po,) = [p for p in load_purchase_orders() if p.message_id == "M15"]

    assert po.customer_code is None
    assert _workbook(rendered, "M15")["PO"]["A1"].value == (po.buyer_name or "").upper()


def test_a_line_is_described_one_way_only() -> None:
    with pytest.raises(ValidationError, match="not both"):
        OrderLine.model_validate(
            {
                "no": 1,
                "customer_item_code": "X-1",
                "attributes": {"colour": "black"},
                "description": "BLACK",
                "quantity": "1",
                "unit_price": "1",
                "requested_date": "2026-12-01",
            }
        )
