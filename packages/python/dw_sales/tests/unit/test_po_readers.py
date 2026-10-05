"""The PO readers: every value with its anchor, every hidden value flagged, every refusal.

The mock attachments are read as committed. Each other case starts from one of
them with a single edit, or from a small file built here, so a test fails for
the one reason it names. The caps are the shipped policy's, narrowed where a
test is about a cap.
"""

from __future__ import annotations

import hashlib
import io
import zipfile
from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from dw_sales.adapters.mock.fixtures import ATTACHMENTS_DIR
from dw_sales.adapters.mock.generate_attachments import FONT_PATH
from dw_sales.adapters.order_rules import load_order_rules
from dw_sales.adapters.readers import (
    ExcelLayout,
    ExcelPoReader,
    PdfLayout,
    PdfPoReader,
    mock_po_readers,
    sniff,
)
from dw_sales.adapters.readers.layouts import (
    MOCK_EXCEL_LAYOUTS,
    MOCK_PDF_LAYOUTS,
    ONE_SHEET,
    VIETNAMESE_PDF,
)
from dw_sales.application.order_ports import PoDocumentReaderPort
from dw_sales.domain.anchors import SourceAnchor
from dw_sales.domain.messages import Attachment, AttachmentContent
from dw_sales.domain.order_checks import IntakeCaps
from dw_sales.domain.orders import PoDocument, RegionFlag, Unreadable

pytestmark = pytest.mark.unit

CAPS = load_order_rules(
    Path(__file__).resolve().parents[5] / "configs" / "policies" / "sales_order_rules@1.0.0.yaml"
).intake
EXCEL = ExcelPoReader(MOCK_EXCEL_LAYOUTS)
PDF = PdfPoReader(MOCK_PDF_LAYOUTS)

M01 = "M01_VLX-PO-2609-0118.xlsx"
M02 = "M02_CVG-260922-07.xlsx"
M03 = "M03_BRN-PO-2609-031.pdf"


def _content(data: bytes) -> AttachmentContent:
    # The declared type is the sender's claim and decides nothing here.
    attachment = Attachment(
        attachment_id="A1",
        name="po.bin",
        media_type="application/octet-stream",
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
    )
    return AttachmentContent(attachment=attachment, data=data)


def _caps(**narrowed: int) -> IntakeCaps:
    return CAPS.model_copy(update=narrowed)


def _fixture(name: str) -> bytes:
    return (ATTACHMENTS_DIR / name).read_bytes()


def _edited(name: str, edit: Callable[[Workbook], object]) -> bytes:
    workbook = load_workbook(io.BytesIO(_fixture(name)))
    edit(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read(data: bytes, customer: str = "VLX", caps: IntakeCaps = CAPS) -> PoDocument | Unreadable:
    return EXCEL.read(_content(data), customer, caps)


def _po(result: PoDocument | Unreadable) -> PoDocument:
    assert isinstance(result, PoDocument), result
    return result


def _refused(result: PoDocument | Unreadable) -> Unreadable:
    assert isinstance(result, Unreadable), result
    return result


def _ref(anchor: SourceAnchor | None) -> str | None:
    return anchor.cell_ref if anchor is not None else None


# ------------------------------------------------------------------- Excel --


def test_an_excel_po_is_read_with_every_value_at_its_cell() -> None:
    data = _fixture(M01)
    po = _po(_read(data))

    header = po.header
    assert (header.po_no, header.revision, header.po_date, header.currency) == (
        "VLX-PO-2609-0118",
        0,
        date(2026, 9, 19),
        "USD",
    )
    assert [_ref(a) for _, a in header.anchors] == ["PO!B4", "PO!E4", "PO!B5", "PO!E5"]
    assert [line.line_no for line in po.lines] == [1, 2, 3, 4]
    third = po.lines[2]
    assert third.model_dump(exclude={"anchors", "flags"}) == {
        "line_no": 3,
        "customer_item_code": "VX-C2F05",
        "description": "MULTI-CORE CABLE 2C 0.5mm2 STRANDED TINNED CU FOIL SHIELD BLACK"
        " REEL 300M SPEC SP-5201",
        "quantity": Decimal(6000),
        "uom": "m",
        "unit_price": Decimal("0.485"),
        "amount": Decimal(2910),
        "requested_date": date(2026, 11, 27),
    }
    assert [_ref(a) for _, a in third.anchors] == [
        "PO!A12",
        "PO!B12",
        "PO!C12",
        "PO!D12",
        "PO!E12",
        "PO!F12",
        "PO!G12",
        "PO!H12",
    ]
    assert third.flags == ()


def test_every_anchor_names_the_file_by_its_sha256_and_the_reader_by_its_version() -> None:
    data = _fixture(M01)
    po = _po(_read(data))

    digest = hashlib.sha256(data).hexdigest()
    assert po.attachment_sha256 == digest
    assert {(a.attachment_id, a.attachment_sha256) for a in po.anchors()} == {("A1", digest)}
    assert po.parser_version == "excel_po_reader@1.1.0"


def test_the_coverage_facts_are_read_with_the_lines() -> None:
    """Rows printed, regions read, the printed total and the buyer named in A1."""
    po = _po(_read(_fixture(M01)))

    assert (po.rows_printed, po.regions, po.unchecked_regions) == (4, ("PO",), ())
    assert po.total is not None
    assert (po.total.amount, _ref(po.total.anchor)) == (Decimal("5814.6"), "PO!G14")
    assert (po.buyer, _ref(po.buyer_anchor)) == ("VELATRIX ELECTRONICS VIETNAM CO., LTD.", "PO!A1")


def test_a_po_printed_one_sheet_per_page_is_read_as_one_po() -> None:
    po = _po(_read(_fixture(M02), "CVG"))

    assert [line.line_no for line in po.lines] == list(range(1, 19))
    assert po.regions == ("Page 1", "Page 2", "Page 3", "Page 4")
    assert _ref(po.header.anchors.po_no) == "Page 1!B4"
    assert _ref(po.lines[5].anchors.line_no) == "Page 2!A10"
    assert _ref(po.lines[17].anchors.unit_price) == "Page 4!F12"
    assert po.lines[10].customer_item_code == "CVG-0006"


def test_pages_are_read_in_page_order_whatever_the_sheet_order() -> None:
    def reverse(workbook: Workbook) -> None:
        for position, title in enumerate(["Page 4", "Page 3", "Page 2", "Page 1"]):
            sheet = workbook[title]
            workbook.move_sheet(sheet, offset=position - workbook.index(sheet))

    in_order = _po(_read(_fixture(M02), "CVG"))
    reversed_sheets = _po(_read(_edited(M02, reverse), "CVG"))

    assert [line.model_dump(exclude={"anchors"}) for line in reversed_sheets.lines] == [
        line.model_dump(exclude={"anchors"}) for line in in_order.lines
    ]


@pytest.mark.parametrize(
    ("edit", "reason", "anchor"),
    [
        (lambda wb: wb.remove(wb["Page 3"]), "pages [1, 2, 4] of [4]", None),
        (lambda wb: wb["Page 3"].__setitem__("E6", "2 / 4"), "page 2 appears twice", None),
        (
            lambda wb: wb["Page 2"].__setitem__("B4", "CVG-260922-08"),
            "shows another PO header",
            "Page 2!B4",
        ),
        (lambda wb: wb["Page 4"].__setitem__("E6", "four"), "is not 'n / N'", "Page 4!E6"),
    ],
    ids=["missing page", "page twice", "another PO", "page number unreadable"],
)
def test_a_paged_po_with_a_page_wrong_is_refused(
    edit: Callable[[Workbook], object], reason: str, anchor: str | None
) -> None:
    refused = _refused(_read(_edited(M02, edit), "CVG"))

    assert reason in refused.reason
    if anchor is not None:
        assert _ref(refused.anchor) == anchor


def test_merged_header_cells_are_read_past_their_merge() -> None:
    """A label merged across two cells has its value right of the merge; a
    title merged over two rows starts the table below both; a value merged over
    two lines belongs to both, anchored where it is printed."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "PO"
    cells: dict[str, str | int | float | datetime] = {
        "A4": "PO No.",
        "C4": "MRG-001",
        "D4": "Revision:",
        "E4": 2,
        "A5": "PO Date",
        "C5": datetime(2026, 9, 19),
        "D5": "Currency",
        "E5": "usd",
        "A9": "No.",
        "B9": "Customer Part No.",
        "C9": "Description",
        "D9": "Quantity",
        "E9": "UoM",
        "F9": "Unit Price",
        "F10": "USD",
        "G10": "per m",
        "I9": "Amount",
        "H9": "Requested Date",
        "A11": 1,
        "B11": "VX-W24-BK",
        "D11": 6100,
        "E11": "M",
        "F11": 0.042,
        "I11": 256.2,
        "H11": datetime(2026, 11, 6),
        "A12": 2,
        "B12": 100234,
        "C12": "  HOOK-UP WIRE\nRED ",
        "D12": 12200,
        "E12": "m",
        "F12": "0.0420",
        "I12": 512.4,
    }
    for cell, value in cells.items():
        sheet[cell] = value
    for merged in ("A4:B4", "A5:B5", "A9:A10", "B9:B10", "C9:C10", "D9:D10", "E9:E10"):
        sheet.merge_cells(merged)
    sheet.merge_cells("F9:G9")
    sheet.merge_cells("H9:H10")
    sheet.merge_cells("I9:I10")
    sheet.merge_cells("H11:H12")
    buffer = io.BytesIO()
    workbook.save(buffer)

    po = _po(_read(buffer.getvalue()))

    assert (po.header.po_no, po.header.revision, po.header.currency) == ("MRG-001", 2, "USD")
    assert [_ref(a) for _, a in po.header.anchors] == ["PO!C4", "PO!E4", "PO!C5", "PO!E5"]
    first, second = po.lines
    assert (first.uom, first.unit_price, first.description) == ("M", Decimal("0.042"), "")
    assert (second.customer_item_code, second.description) == ("100234", "HOOK-UP WIRE RED")
    assert second.requested_date == date(2026, 11, 6)
    assert _ref(second.anchors.requested_date) == "PO!H11"
    assert _ref(second.anchors.unit_price) == "PO!F12"
    assert po.total is None


def test_a_float_reads_as_the_number_printed() -> None:
    def noisy(workbook: Workbook) -> None:
        workbook["PO"]["F10"] = 0.04200000000000001

    assert _po(_read(_edited(M01, noisy))).lines[0].unit_price == Decimal("0.042")


def test_a_unit_dw1_has_no_word_for_is_read_as_printed() -> None:
    """`uom_mismatch` is the check's to raise on the line, not a file refused."""
    po = _po(_read(_edited(M01, lambda wb: wb["PO"].__setitem__("E11", "roll"))))

    assert po.lines[1].uom == "roll"


# --------------------------------------------------- what a person can't see --


def _hide_sheet(workbook: Workbook) -> None:
    workbook.create_sheet("Cover")  # a workbook needs one visible sheet to save
    workbook["PO"].sheet_state = "hidden"


def _font_like_fill(workbook: Workbook) -> None:
    cell = workbook["PO"]["D11"]
    cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    cell.font = Font(color="FFFFFFFF")


@pytest.mark.parametrize(
    ("edit", "line", "fields", "flag"),
    # line None: every line (a hidden column hides one value of each).
    [
        (
            lambda wb: setattr(wb["PO"].row_dimensions[11], "hidden", True),
            2,
            {
                "line_no",
                "customer_item_code",
                "description",
                "quantity",
                "uom",
                "unit_price",
                "amount",
                "requested_date",
            },
            RegionFlag.HIDDEN_ROW,
        ),
        (
            lambda wb: setattr(wb["PO"].column_dimensions["D"], "hidden", True),
            None,
            {"quantity"},
            RegionFlag.HIDDEN_COLUMN,
        ),
        (_font_like_fill, 2, {"quantity"}, RegionFlag.FONT_MATCHES_FILL),
        (
            lambda wb: wb["PO"].__setitem__("C12", "=CONCATENATE(B12)"),
            3,
            {"description"},
            RegionFlag.FORMULA_WITHOUT_CACHED_VALUE,
        ),
    ],
    ids=["hidden row", "hidden column", "font matches fill", "formula without cached value"],
)
def test_a_value_a_person_would_not_see_is_read_and_flagged(
    edit: Callable[[Workbook], object], line: int | None, fields: set[str], flag: RegionFlag
) -> None:
    """Read, never skipped: the lines read stay the lines the file holds, and
    each flagged value raises `value_uncertain` (test_order_checks)."""
    po = _po(_read(_edited(M01, edit)))

    assert [x.line_no for x in po.lines] == [1, 2, 3, 4]
    for x in po.lines:
        expected = fields if line in (None, x.line_no) else set()
        assert {(f.field, f.flag) for f in x.flags} == {(name, flag) for name in expected}


def test_a_hidden_sheet_is_read_with_every_value_flagged() -> None:
    po = _po(_read(_edited(M01, _hide_sheet)))

    assert {f.flag for f in po.header.flags} == {RegionFlag.HIDDEN_SHEET}
    assert all({f.flag for f in line.flags} == {RegionFlag.HIDDEN_SHEET} for line in po.lines)
    assert po.total is not None and {f.flag for f in po.total.flags} == {RegionFlag.HIDDEN_SHEET}


def test_the_mock_po_with_a_hidden_row_and_a_hidden_page_flags_exactly_those_lines() -> None:
    po = _po(_read(_fixture("M27_CVG-261001-02.xlsx"), "CVG"))

    flagged = {line.line_no: {f.flag for f in line.flags} for line in po.lines if line.flags}
    assert flagged == {
        2: {RegionFlag.HIDDEN_ROW},
        4: {RegionFlag.HIDDEN_SHEET},
        5: {RegionFlag.HIDDEN_SHEET},
        6: {RegionFlag.HIDDEN_SHEET},
    }
    assert po.header.flags == ()
    assert sum(line.amount for line in po.lines) == po.total.amount if po.total else False


def test_a_sheet_no_layout_reads_is_reported_as_unchecked() -> None:
    def notes(workbook: Workbook) -> None:
        workbook.create_sheet("Notes")["A1"] = "Second delivery address"

    po = _po(_read(_edited(M01, notes)))

    assert po.unchecked_regions == ("Notes",)


def test_a_numbered_row_below_a_gap_is_counted_as_printed_and_not_read() -> None:
    """The lines under the gap are printed; `line_total_mismatch` reports them."""
    po = _po(_read(_edited(M01, lambda wb: wb["PO"].insert_rows(12))))

    assert [line.line_no for line in po.lines] == [1, 2]
    assert po.rows_printed == 4


def test_line_numbers_are_read_as_printed_even_when_they_skip() -> None:
    po = _po(_read(_fixture("M22_CVG-260930-12.xlsx"), "CVG"))

    assert [line.line_no for line in po.lines] == [1, 2, 3, 4, 6, 7, 8]


# ---------------------------------------------------------------- refusals --


@pytest.mark.parametrize(
    ("cell", "value", "reason"),
    [
        ("D11", "many", "row 11: quantity is not a number"),
        ("H11", "next month", "row 11: requested date is not a date"),
        ("D11", 0, "quantity"),
        ("B11", None, "customer_item_code"),
    ],
)
def test_a_line_value_that_cannot_be_read_names_its_cell_and_not_its_value(
    cell: str, value: str | int | None, reason: str
) -> None:
    def put(workbook: Workbook) -> None:
        workbook["PO"][cell] = value

    refused = _refused(_read(_edited(M01, put)))

    assert reason in refused.reason
    assert _ref(refused.anchor) == f"PO!{cell}"
    if isinstance(value, str):
        assert value not in refused.reason


@pytest.mark.parametrize(
    ("cell", "value", "reason"),
    [("E4", None, "revision is not a whole number"), ("E5", "EUR", "currency")],
)
def test_a_header_value_that_cannot_be_read_names_its_cell(
    cell: str, value: str | None, reason: str
) -> None:
    def put(workbook: Workbook) -> None:
        workbook["PO"][cell] = value

    refused = _refused(_read(_edited(M01, put)))

    assert reason in refused.reason
    assert _ref(refused.anchor) == f"PO!{cell}"
    assert "EUR" not in refused.reason


def test_a_label_printed_twice_is_refused_rather_than_picked() -> None:
    refused = _refused(_read(_edited(M01, lambda wb: wb["PO"].__setitem__("G2", "PO No."))))

    assert "a label of the layout appears 2 times" in refused.reason


def test_two_sheets_with_a_po_table_are_refused_for_a_one_sheet_customer() -> None:
    refused = _refused(_read(_edited(M01, lambda wb: wb.copy_worksheet(wb["PO"]))))

    assert refused.reason == "2 sheets have a PO table: PO, PO Copy"


def test_a_request_for_quotation_is_not_a_po() -> None:
    refused = _refused(_read(_fixture("M09_QRL-RFQ-2609-03.xlsx"), "QRL"))

    assert refused.reason == "no sheet has the PO table"


def test_a_customer_without_a_layout_is_not_read() -> None:
    assert _refused(_read(_fixture(M01), "BRN")).reason == "no Excel PO layout for customer BRN"


def test_a_file_that_is_not_a_workbook_is_refused() -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.writestr("notes.txt", "not a workbook")

    assert _refused(_read(_fixture(M03))).reason == "not an .xlsx workbook"
    assert "openpyxl can open" in _refused(_read(archive.getvalue())).reason


def test_a_macro_enabled_workbook_is_refused() -> None:
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(_fixture(M01))) as source,
        zipfile.ZipFile(buffer, "w") as target,
    ):
        for entry in source.infolist():
            target.writestr(entry, source.read(entry))
        target.writestr("xl/vbaProject.bin", b"\x00" * 16)

    assert _refused(_read(buffer.getvalue())).reason == "a macro-enabled workbook is not read"


@pytest.mark.parametrize(
    ("caps", "reason"),
    [
        (_caps(max_attachment_bytes=1_000), "over the 1000 bytes read here"),
        (_caps(max_unpacked_bytes=1_000), "unpacks past the 1000 bytes"),
        (_caps(max_sheets=3), "more than the 3 sheets read here"),
        (_caps(max_cells=100), "more than the 100 cells read here"),
    ],
    ids=["bytes", "unpacked", "sheets", "cells"],
)
def test_a_workbook_past_a_policy_cap_is_refused(caps: IntakeCaps, reason: str) -> None:
    refused = _refused(_read(_fixture(M02), "CVG", caps))

    assert reason in refused.reason


def test_an_excel_layout_names_every_value() -> None:
    columns = dict(ONE_SHEET.column_labels)
    del columns["amount"]

    with pytest.raises(ValueError, match="every line field"):
        ExcelLayout(header_labels=ONE_SHEET.header_labels, column_labels=columns)


def test_the_buyer_is_read_without_a_customer_layout() -> None:
    """How a forwarded PO is attributed before anyone knows whose layout it is."""
    buyer = EXCEL.buyer(_content(_fixture("M15_TVT-PO-0929-01.xlsx")), CAPS)

    assert buyer is not None
    assert (buyer.name, buyer.anchor.cell_ref) == ("TALVERRA TRADING CO., LTD.", "PO!A1")
    assert EXCEL.buyer(_content(b"not a workbook"), CAPS) is None


# --------------------------------------------------------------------- PDF --


def _read_pdf(
    data: bytes, customer: str = "BRN", caps: IntakeCaps = CAPS
) -> PoDocument | Unreadable:
    return PDF.read(_content(data), customer, caps)


# Where the mock customer prints each column, and whether it is right-aligned.
_PDF_COLUMNS = (
    (55, True),
    (62, False),
    (140, False),
    (612, True),
    (618, False),
    (690, True),
    (770, True),
    (778, False),
)
_TITLES = (
    "STT",
    "Mã hàng KH",
    "Mô tả hàng hóa",
    "Số lượng",
    "ĐVT",
    "Đơn giá",
    "Thành tiền",
    "Ngày giao",
)


def _vietnamese_pdf(*rows: tuple[str, ...] | str, header: bool = True) -> bytes:
    """A one-page PO in the Vietnamese layout: each row is its eight cells, or a
    string printed under the description column."""
    pdfmetrics.registerFont(TTFont("ReaderTest", str(FONT_PATH)))
    buffer = io.BytesIO()
    width, height = landscape(A4)
    canvas = Canvas(buffer, pagesize=(width, height), invariant=1)
    y = height - 50
    if header:
        canvas.setFont("ReaderTest", 9)
        for text in (
            "Số đơn hàng: BRN-PO-2610-001",
            "Lần sửa đổi: 0",
            "Ngày đặt hàng: 01/10/2026",
            "Đơn vị tiền tệ: VND",
        ):
            y -= 15
            canvas.drawString(40, y, text)
    for row in (_TITLES, *rows):
        y -= 14
        canvas.setFont("ReaderTest", 7.5)
        if isinstance(row, str):
            canvas.drawString(140, y, row)
            continue
        for (x, right), value in zip(_PDF_COLUMNS, row, strict=True):
            (canvas.drawRightString if right else canvas.drawString)(x, y, value)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()


_ROW = (
    "1",
    "BRN-W-0007",
    "HOOK-UP WIRE 1C AWG24",
    "6.100",
    "m",
    "1.150",
    "7.015.000",
    "16/11/2026",
)


def test_a_text_pdf_is_read_with_every_value_at_its_page_and_box() -> None:
    po = _po(_read_pdf(_fixture(M03)))

    header = po.header
    assert (header.po_no, header.revision, header.po_date, header.currency) == (
        "BRN-PO-2609-031",
        0,
        date(2026, 9, 23),
        "VND",
    )
    assert [(a.page, a.quote) for _, a in header.anchors] == [
        (1, "BRN-PO-2609-031"),
        (1, "0"),
        (1, "23/09/2026"),
        (1, "VND"),
    ]
    assert [(line.quantity, line.unit_price, line.amount) for line in po.lines] == [
        (Decimal(6100), Decimal(1150), Decimal(7015000)),
        (Decimal(3000), Decimal(14500), Decimal(43500000)),
        (Decimal(2000), Decimal(21800), Decimal(43600000)),
    ]
    second = po.lines[1]
    assert (second.customer_item_code, second.requested_date) == ("BRN-C-0112", date(2026, 11, 30))
    assert second.description == (
        "MULTI-CORE CABLE 2C 0.75mm2 STRANDED BARE CU NO SHIELD GREY REEL 300M SPEC SP-5202"
    )
    assert (second.anchors.quantity.page, second.anchors.quantity.quote) == (1, "3.000")
    assert po.total is not None and po.total.amount == Decimal(94115000)
    assert (po.rows_printed, po.regions, po.parser_version) == (
        3,
        ("page 1",),
        "pdf_po_reader@1.1.0",
    )


def test_a_pdf_value_is_boxed_where_it_is_printed_never_by_a_text_line() -> None:
    """ADR 0009: the box is on the rendered page; each row's boxes sit lower
    than the row above, each cell right of the one before."""
    po = _po(_read_pdf(_fixture(M03)))

    rows = [line.anchors for line in po.lines]
    tops = [anchors.quantity.boxes[0].y for anchors in rows]
    assert tops == sorted(tops) and len(set(tops)) == 3
    lefts = [anchor.boxes[0].x for _, anchor in rows[0]]
    assert lefts == sorted(lefts)
    for anchors in rows:
        for _, anchor in anchors:
            assert anchor.page == 1 and len(anchor.boxes) == 1
            box = anchor.boxes[0]
            assert box.x + box.w <= 1 and box.y + box.h <= 1


def test_a_built_pdf_is_read_like_the_mock_one() -> None:
    """The cases below use PDFs built here; this shows one is read at all."""
    po = _po(_read_pdf(_vietnamese_pdf(_ROW)))

    assert [(line.line_no, line.quantity) for line in po.lines] == [(1, Decimal(6100))]


def test_a_row_outside_the_table_is_counted_as_printed_and_not_read() -> None:
    """A wrapped description ends the table; the row under it is printed and
    not read, which `line_total_mismatch` reports."""
    second = (
        "2",
        "BRN-C-0112",
        "MULTI-CORE CABLE",
        "3.000",
        "m",
        "14.500",
        "43.500.000",
        "30/11/2026",
    )

    po = _po(_read_pdf(_vietnamese_pdf(_ROW, "STRANDED TINNED CU", second)))

    assert ([line.line_no for line in po.lines], po.rows_printed) == ([1], 2)


@pytest.mark.parametrize(
    ("cell", "value", "reason"),
    [
        (3, "6.1", "quantity is not a number"),
        (7, "2026-11-16", "requested date is not a date"),
    ],
)
def test_a_pdf_cell_that_cannot_be_read_names_its_field_and_not_its_value(
    cell: int, value: str, reason: str
) -> None:
    row = list(_ROW)
    row[cell] = value

    refused = _refused(_read_pdf(_vietnamese_pdf(tuple(row))))

    assert reason in refused.reason
    assert value not in refused.reason
    assert refused.anchor is not None and refused.anchor.page == 1


def test_a_pdf_without_a_text_layer_is_a_scan_and_unreadable() -> None:
    refused = _refused(_read_pdf(_fixture("M16_BRN-PO-2609-044.pdf")))

    assert refused.reason == "the PDF has no text layer: a scan or an image"


def test_a_password_protected_pdf_is_unreadable() -> None:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_fixture(M03))))
    writer.encrypt(user_password="secret", algorithm="RC4-128")
    buffer = io.BytesIO()
    writer.write(buffer)

    assert _refused(_read_pdf(buffer.getvalue())).reason == "the PDF is password-protected"


def test_a_pdf_longer_than_the_policy_cap_is_refused_before_it_is_read() -> None:
    refused = _refused(_read_pdf(_fixture(M03), caps=_caps(max_pages=1)))

    assert refused.reason == "the PDF has more than the 1 pages read here"


def test_a_file_that_is_not_a_pdf_and_a_customer_without_a_layout_are_refused() -> None:
    assert _refused(_read_pdf(_fixture(M01))).reason == "not a PDF pypdf can read"
    assert _refused(_read_pdf(_fixture(M03), "VLX")).reason == "no PDF PO layout for customer VLX"


def test_a_pdf_layout_has_one_column_per_value_and_reads_the_number_first() -> None:
    columns = VIETNAMESE_PDF.columns

    with pytest.raises(ValueError, match="one column for each line field"):
        PdfLayout(**{**VIETNAMESE_PDF.__dict__, "columns": columns[:-1]})
    with pytest.raises(ValueError, match="first column"):
        PdfLayout(**{**VIETNAMESE_PDF.__dict__, "columns": (columns[1], columns[0], *columns[2:])})
    with pytest.raises(ValueError, match="separators"):
        PdfLayout(**{**VIETNAMESE_PDF.__dict__, "decimal_separator": "."})


# ----------------------------------------------------------- what bytes are --


def _zip(*names: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name in names:
            archive.writestr(name, b"")
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("data", "kind"),
    [
        (_fixture(M01), "xlsx"),
        (_fixture(M03), "pdf"),
        (_fixture("M16_BRN-PO-2609-044.pdf"), "pdf"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "unsupported"),
        (b"\x89PNG\r\n\x1a\n" + b"\x00" * 16, "unsupported"),
        (b"\xff\xd8\xff\xe0" + b"\x00" * 16, "unsupported"),
        (_zip("word/document.xml"), "unsupported"),
        (b"PK\x03\x04 broken", "unsupported"),
        (b"Dear Sales, here is our order.", None),
    ],
    ids=["xlsx", "pdf", "scan", "ole", "png", "jpeg", "docx", "broken zip", "text"],
)
def test_a_file_is_what_its_bytes_say(data: bytes, kind: str | None) -> None:
    assert sniff(data) == kind


def test_the_readers_satisfy_the_port() -> None:
    port: PoDocumentReaderPort = mock_po_readers()

    assert port.sniff(_fixture(M01)) == "xlsx"


async def test_each_file_goes_to_the_reader_its_bytes_name_whatever_it_is_called() -> None:
    readers = mock_po_readers()

    excel = await readers.read(_content(_fixture(M01)), "VLX", CAPS)
    pdf = await readers.read(_content(_fixture(M03)), "BRN", CAPS)
    image = await readers.read(_content(b"\x89PNG\r\n\x1a\n"), "VLX", CAPS)

    assert _po(excel).header.po_no == "VLX-PO-2609-0118"
    assert _po(pdf).header.po_no == "BRN-PO-2609-031"
    assert _refused(image).reason == "not a PO format read in this slice (text PDF or .xlsx)"


async def test_the_buyer_is_read_only_from_a_workbook() -> None:
    readers = mock_po_readers()

    assert await readers.buyer(_content(_fixture(M03)), CAPS) is None
    buyer = await readers.buyer(_content(_fixture("M13_NRV-PO-26-0471.xlsx")), CAPS)
    assert buyer is not None and buyer.name.startswith("NORVANTA")
