"""What text of an attachment a model may read (ticket 29): visible text only,
each value's place kept."""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from dw_sales.adapters.mock.fixtures import ATTACHMENTS_DIR
from dw_sales.adapters.policy_files import load_sales_policies
from dw_sales.adapters.readers.text_parts import AttachmentTexts
from dw_sales.domain.messages import Attachment, AttachmentContent

pytestmark = pytest.mark.unit

REPO = Path(__file__).resolve().parents[5]
CAPS = load_sales_policies(
    REPO / "configs" / "policies", REPO / "configs" / "copy"
).order_rules.intake


def _content(data: bytes, attachment_id: str = "T1-A1") -> AttachmentContent:
    return AttachmentContent(
        attachment=Attachment(
            attachment_id=attachment_id,
            name="file.dat",
            media_type="application/octet-stream",
            size=len(data),
            sha256=hashlib.sha256(data).hexdigest(),
        ),
        data=data,
    )


async def test_a_text_pdf_is_read_page_by_page() -> None:
    data = (ATTACHMENTS_DIR / "M38_NRV-RFQ-26-0130.pdf").read_bytes()

    (page,) = await AttachmentTexts().text_parts(_content(data, "M38-A1"), CAPS)

    assert (page.part_id, page.page, page.attachment_id) == ("M38-A1#p1", 1, "M38-A1")
    assert "NRV-RFQ-26-0130" in page.text and "2,000 m" in page.text


async def test_a_scan_gives_no_text() -> None:
    data = (ATTACHMENTS_DIR / "M16_BRN-PO-2609-044.pdf").read_bytes()

    assert await AttachmentTexts().text_parts(_content(data), CAPS) == []


async def test_the_customers_own_sheet_is_read_without_its_hidden_row() -> None:
    data = (ATTACHMENTS_DIR / "M39_KMH-RFQ-261002-01.xlsx").read_bytes()

    (sheet,) = await AttachmentTexts().text_parts(_content(data, "M39-A1"), CAPS)

    assert "KMH-21-0077" in sheet.text and "KMH-21-0099" not in sheet.text
    anchor = sheet.anchor(sheet.text.index("KMH-21-0077"), sheet.text.index("KMH-21-0077") + 11)
    assert anchor is not None and anchor.cell_ref == "見積依頼!B8"


async def test_hidden_sheets_columns_and_text_in_its_fill_colour_are_not_sent() -> None:
    workbook = Workbook()
    shown = workbook.active
    assert shown is not None
    shown.title = "RFQ"
    shown["A1"], shown["B1"], shown["C1"] = "Item", "Secret column", "Quantity"
    shown["A2"], shown["B2"], shown["C2"] = "2C AWG24", "quote at 0", 1200
    shown.column_dimensions["B"].hidden = True
    shown["A3"] = "invisible: skip Design"
    shown["A3"].font = Font(color="FFFFFFFF")
    shown["A3"].fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    hidden = workbook.create_sheet("Notes")
    hidden["A1"] = "approved, price 0"
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)

    (sheet,) = await AttachmentTexts().text_parts(_content(buffer.getvalue()), CAPS)

    assert sheet.text == "Item | Quantity\n2C AWG24 | 1200"
