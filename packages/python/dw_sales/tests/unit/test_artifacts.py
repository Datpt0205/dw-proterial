"""Ticket 06: every artifact rendered from the case, stored, and served only in
its state and for the version it was rendered at.

The cases are the mock mailbox's own, walked through the domain's moves; the
service runs over an in-memory unit of work, the shipped copy and policies,
and the real file writer. Per kind: rendered at its state with its template
``id@version``, case version and sha256; refused before its state; refused
after an edit that bumped the case version; and the bytes served are the
bytes recorded.
"""

from __future__ import annotations

import email
import hashlib
import io
import uuid
import zipfile
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from email import policy
from pathlib import Path
from types import TracebackType
from typing import Any, Self, cast

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from dw_kernel.errors import (
    ConflictError,
    DomainError,
    InfrastructureError,
    NotFoundError,
    PermissionDeniedError,
)
from dw_kernel.ids import TenantId, WorkspaceId
from dw_kernel.ports import FixedClock, Uuid4Generator
from dw_platform.application.access_context import AccessContext
from dw_platform.application.authorization import ScopeAuthorizationService
from dw_platform.application.directory import WorkspaceMember
from dw_platform.domain.audit import AuditEvent
from dw_sales.adapters.artifact_files import ArtifactFiles
from dw_sales.adapters.memory_rules import MemoryRuleVersions, platform_book
from dw_sales.adapters.mock import MockInbox, MockSalesCatalog
from dw_sales.adapters.order_rules import PlatformOrderRules, load_order_rules
from dw_sales.adapters.policy_files import (
    load_bravo_upload,
    load_copy,
    load_pricing,
    load_quote_rules,
)
from dw_sales.adapters.reader_client import (
    ClientDesignReplyReader,
    ClientPoReaders,
    ClientQuoteForms,
    ClientRfqReader,
)
from dw_sales.adapters.readers.operations import InProcessReaderTransport
from dw_sales.application.access import Gate, PriceView
from dw_sales.application.artifact_content import (
    EML,
    PDF,
    XLSX,
    SalesDocumentCopy,
    SalesEmailCopy,
)
from dw_sales.application.artifacts_service import (
    ORDER_GATES,
    QUOTE_GATES,
    ArtifactService,
    ArtifactView,
)
from dw_sales.application.case_store import ArtifactRecord, CaseOrigin, Stored
from dw_sales.application.drafting import ArtifactCopy, ArtifactCopyPort, Drafting, People
from dw_sales.application.order_artifacts import portal_checklist
from dw_sales.application.order_intake import OrderIntake
from dw_sales.application.ports import SalesScope
from dw_sales.application.quotation import QuotationService, QuoteRules, ReplyAttached
from dw_sales.application.rules import Policy, TenantRuleBook
from dw_sales.application.views import order_line
from dw_sales.domain.catalog import LmeBand, LmeMonth, Quotation
from dw_sales.domain.dispositions import CaseKind
from dw_sales.domain.orders import (
    Accepted,
    Actor,
    AskCustomer,
    Capability,
    CloseReason,
    CorrectedBySales,
    FindingCode,
    MappingStatus,
    OrderCase,
    OrderStatus,
)
from dw_sales.domain.quotes import (
    DeclineReason,
    LinePrice,
    PricingDecision,
    QuoteActor,
    QuoteCapability,
    QuoteCase,
)
from dw_sales.evals.world import MemoryStore

pytestmark = pytest.mark.unit

REPO = Path(__file__).resolve().parents[5]
POLICIES = REPO / "configs" / "policies"
COPY = REPO / "configs" / "copy"
SCOPE = SalesScope(TenantId(uuid.UUID(int=1)), WorkspaceId(uuid.UUID(int=2)))
OTHER = SalesScope(TenantId(uuid.UUID(int=9)), WorkspaceId(uuid.UUID(int=10)))

AN = uuid.UUID(int=0xA1)  # PIC đơn hàng, export control
DIEU = uuid.UUID(int=0xD1)  # PIC báo giá, cross-checks An's orders
GIANG = uuid.UUID(int=0x61)  # head, approves
MEMBERS = (
    WorkspaceMember(AN, "Nguyễn Văn An", "an.nguyen@alpha.local", ("sales_pic",), (), "sales"),
    WorkspaceMember(DIEU, "Hoàng Thị Diệu", "dieu.hoang@alpha.local", ("sales_pic",), (), "sales"),
    WorkspaceMember(GIANG, "Đỗ Minh Giang", "giang.do@alpha.local", ("sales_head",), (), "sales"),
)
T0 = datetime(2026, 10, 2, 2, 0, tzinfo=UTC)
NOW = datetime(2026, 10, 5, 3, 0, tzinfo=UTC)
CONFIRMED_ON = date(2026, 11, 27)
YCBG_M10 = "YCBG-2609-030"
QUOTE_NO = "Q26-0301"
BAND = LmeBand(low_usd_per_tonne=Decimal(10500), high_usd_per_tonne=Decimal(11000))
DECIDED = Decimal("0.6890")
# M10's evidence only Sales may see: other customers' prices and the reference.
INTERNAL_PRICES = ("0.7120", "0.6980", "0.7050")
GUIDANCE = "Chỉ đạo nội bộ: không thấp hơn giá khách khác"

QUOTE_RULES = load_quote_rules(POLICIES / "sales_quote_rules@1.1.0.yaml")
PRICING = load_pricing(POLICIES / "sales_pricing@1.0.0.yaml")
ARTIFACT_COPY = ArtifactCopy(
    emails=load_copy(COPY / "sales_emails@1.2.0.yaml", SalesEmailCopy),
    documents=load_copy(COPY / "sales_documents@1.2.0.yaml", SalesDocumentCopy),
    upload=load_bravo_upload(POLICIES / "sales_bravo_upload@1.1.0.yaml"),
    design_mailboxes=QUOTE_RULES.design_mailboxes,
)
# The platform layer every service here resolves through: no tenant holds a
# version of its own until a test proposes one.
BOOK = platform_book(
    QUOTE_RULES, PRICING, ARTIFACT_COPY.emails, ARTIFACT_COPY.documents, ARTIFACT_COPY.upload
)
WRITER = ArtifactFiles()
GATE = Gate(ScopeAuthorizationService())

PIC_SCOPES = frozenset(
    {
        "sales.case.read",
        "sales.price.read",
        "sales.order.prepare",
        "sales.order.cross_check",
        "sales.quote.prepare",
    }
)


def context(
    principal: uuid.UUID = AN, scopes: frozenset[str] = PIC_SCOPES, scope: SalesScope = SCOPE
) -> AccessContext:
    return AccessContext(
        tenant_id=scope.tenant_id.value,
        workspace_id=scope.workspace_id.value,
        principal_id=principal,
        roles=frozenset({"sales_pic"}),
        scopes=scopes,
        plan_id="professional",
    )


# ------------------------------------------------------- the in-memory world --


@dataclass
class World:
    """The stores of every scope: cases, artifact records, objects, audit rows."""

    orders: dict[tuple[SalesScope, uuid.UUID], OrderCase] = field(default_factory=dict)
    quotes: dict[tuple[SalesScope, uuid.UUID], QuoteCase] = field(default_factory=dict)
    artifacts: dict[tuple[SalesScope, uuid.UUID], ArtifactRecord] = field(default_factory=dict)
    objects: dict[str, bytes] = field(default_factory=dict)
    audit: list[AuditEvent] = field(default_factory=list)
    commits: int = 0
    # Run once, as the next object is stored: another request in between.
    meanwhile: Callable[[], Awaitable[None]] | None = None
    # The next commit fails: before anything is kept, or after (a commit that
    # went through and failed on the way back).
    commit_fails: str | None = None

    def put(self, case: OrderCase | QuoteCase, scope: SalesScope = SCOPE) -> None:
        if isinstance(case, OrderCase):
            self.orders[(scope, case.case_id)] = case
        else:
            self.quotes[(scope, case.case_id)] = case


class _Store[CaseT]:
    def __init__(self, held: dict[tuple[SalesScope, uuid.UUID], CaseT], scope: SalesScope):
        self.held, self.scope = held, scope

    async def get(self, case_id: uuid.UUID) -> Stored[CaseT] | None:
        case = self.held.get((self.scope, case_id))
        return None if case is None else Stored(case, CaseOrigin(None, None))


class _Artifacts:
    def __init__(self, world: World, scope: SalesScope) -> None:
        self.world, self.scope = world, scope
        self.pending: list[ArtifactRecord] = []

    async def add(self, artifact: ArtifactRecord) -> bool:
        """As the store's unique key does: one file of a kind and content
        type per case version, among what is committed and what is pending."""
        key = _render_key(artifact)
        held = [*await self.for_case(artifact.case_kind, artifact.case_id), *self.pending]
        if any(_render_key(record) == key for record in held):
            return False
        self.pending.append(artifact)
        return True

    async def get(self, artifact_id: uuid.UUID) -> ArtifactRecord | None:
        return self.world.artifacts.get((self.scope, artifact_id))

    async def for_case(self, case_kind: CaseKind, case_id: uuid.UUID) -> list[ArtifactRecord]:
        return [
            r
            for (scope, _), r in self.world.artifacts.items()
            if scope == self.scope and (r.case_kind, r.case_id) == (case_kind, case_id)
        ]


def _render_key(record: ArtifactRecord) -> tuple[object, ...]:
    return (record.case_kind, record.case_id, record.case_version, record.kind, record.content_type)


class _Audit:
    def __init__(self) -> None:
        self.pending: list[AuditEvent] = []

    async def append(self, event: AuditEvent) -> None:
        self.pending.append(event)


class Work:
    """One transaction: what it added is kept only when it commits."""

    def __init__(self, world: World, scope: SalesScope) -> None:
        self.world, self.scope = world, scope
        self.orders = _Store(world.orders, scope)
        self.quotes = _Store(world.quotes, scope)
        self.artifacts = _Artifacts(world, scope)
        self.audit = _Audit()

    async def commit(self) -> None:
        fails, self.world.commit_fails = self.world.commit_fails, None
        if fails == "before":
            raise InfrastructureError("the database went away")
        for record in self.artifacts.pending:
            self.world.artifacts[(self.scope, record.artifact_id)] = record
        self.world.audit += self.audit.pending
        self.world.commits += 1
        if fails == "after":
            raise InfrastructureError("the commit's answer was lost")

    async def __aenter__(self) -> Self:
        return self

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        tb: TracebackType | None,
    ) -> None:
        return None


class Objects:
    def __init__(self, world: World) -> None:
        self.world = world

    async def put_object(self, key: str, data: bytes, content_type: str) -> str:
        self.world.objects[key] = data
        meanwhile, self.world.meanwhile = self.world.meanwhile, None
        if meanwhile is not None:
            await meanwhile()
        return key

    async def get_object(self, key: str) -> bytes:
        return self.world.objects[key]

    async def delete_object(self, key: str) -> None:
        del self.world.objects[key]


class Directory:
    async def list_members(self, context: AccessContext) -> list[WorkspaceMember]:
        return list(MEMBERS)


class QuoteCases:
    def __init__(self) -> None:
        self.cases: list[QuoteCase] = []

    async def cases_for_ycbg(self, scope: SalesScope, ycbg_no: str) -> Sequence[QuoteCase]:
        return [c for c in self.cases if c.ycbg is not None and c.ycbg.ycbg_no == ycbg_no]


class OrderCases:
    def __init__(self) -> None:
        self.cases: dict[uuid.UUID, OrderCase] = {}

    async def cases_for_po(
        self, scope: SalesScope, customer_code: str, po_no: str
    ) -> Sequence[OrderCase]:
        return [
            c
            for c in self.cases.values()
            if (c.customer_code, c.header.po_no) == (customer_code, po_no)
        ]


def quotation_service(cases: QuoteCases | None = None) -> QuotationService:
    catalog = MockSalesCatalog.load(SCOPE)
    return QuotationService(
        catalog=catalog,
        inbox=MockInbox.load(SCOPE),
        rfq_reader=ClientRfqReader(InProcessReaderTransport()),
        reply_reader=ClientDesignReplyReader(InProcessReaderTransport()),
        cases=cases or QuoteCases(),
        ledger=catalog,
        policies=BOOK,
    )


def _factory(world: World) -> Any:
    """A unit of work per scope over ``world``: as much of the port as the
    service uses."""
    return lambda scope: Work(world, scope)


def service(world: World, copy: ArtifactCopyPort = BOOK) -> ArtifactService:
    return ArtifactService(
        uow=_factory(world),
        gate=GATE,
        storage=Objects(world),
        clock=FixedClock(NOW),
        ids=Uuid4Generator(),
        writer=WRITER,
        copy=copy,
        catalog=MockSalesCatalog.load(SCOPE),
        directory=Directory(),
        design=quotation_service(),
    )


# ------------------------------------------------------------ order cases --

_INTAKE_RULES = PlatformOrderRules(load_order_rules(POLICIES / "sales_order_rules@1.5.0.yaml"))


def intake(cases: OrderCases) -> OrderIntake:
    return OrderIntake(
        catalog=MockSalesCatalog.load(SCOPE),
        inbox=MockInbox.load(SCOPE),
        reader=ClientPoReaders(InProcessReaderTransport()),
        quote_forms=ClientQuoteForms(InProcessReaderTransport()),
        rules=_INTAKE_RULES,
        cases=cases,
        new_case_id=uuid.uuid4,
    )


async def in_review(message_id: str, cases: OrderCases | None = None) -> OrderCase:
    """The order a message opened, handed to the PIC as DW1 does at intake."""
    cases = cases if cases is not None else OrderCases()
    message = await MockInbox.load(SCOPE).get_message(SCOPE, message_id)
    assert message is not None
    outcome = await intake(cases).process(SCOPE, message)
    assert outcome.case is not None, message_id
    case = outcome.case.start_review()
    cases.cases[case.case_id] = case
    return case


def decided_all(case: OrderCase, by: uuid.UUID = AN) -> OrderCase:
    """Every open finding decided the way its code allows, by ``by``. A
    `value_uncertain` is typed field by field first (`typed_all`)."""
    actor = Actor(user_id=by, capabilities=frozenset({Capability.ACKNOWLEDGE_EXPORT_CONTROL}))
    for finding in case.findings:
        if not finding.is_open:
            continue
        if finding.code is FindingCode.CUSTOMER_UNKNOWN:
            value = case.customer_code
            disposition: Any = CorrectedBySales(value=value, source="PO gốc", by=by, at=T0)
        else:
            disposition = Accepted(reason="đã xem", by=by, at=T0)
        case = case.dispose(finding.key, disposition, actor)
    return case


# M27 line 2's price as Sales reads it on the hidden row: not the price read.
CORRECTED_PRICE = Decimal("0.0475")


async def typed_all(case: OrderCase, by: uuid.UUID = AN) -> OrderCase:
    """Every value each `value_uncertain` flagged, typed by ``by`` as read,
    except line 2's unit price, typed as `CORRECTED_PRICE`."""
    actor = Actor(user_id=by)
    for finding in case.findings:
        if finding.code is not FindingCode.VALUE_UNCERTAIN:
            continue
        line = case.line(finding.line_no) if finding.line_no is not None else None
        for name in sorted(case.flagged_fields(finding)):
            assert line is not None
            value = str(getattr(line.po_line, name))
            if (line.line_no, name) == (2, "unit_price"):
                value = str(CORRECTED_PRICE)
            case = await intake(OrderCases()).correct_value(
                SCOPE, case, finding.key, name, value, actor, T0
            )
    return case


def prepared(case: OrderCase) -> OrderCase:
    return case.prepare(Actor(user_id=AN), T0)


def uploaded(case: OrderCase) -> OrderCase:
    return prepared(case).record_bravo_entry(
        "SO26-1001", Actor(user_id=AN), T0, entry_compared=True
    )


def cross_checked(case: OrderCase) -> OrderCase:
    return uploaded(case).cross_check(Actor(user_id=DIEU), T0)


def confirmed(case: OrderCase) -> OrderCase:
    checked = cross_checked(case)
    return checked.confirm(
        Actor(user_id=AN),
        T0,
        {line.line_no: [CONFIRMED_ON] * line.deliveries for line in checked.lines},
    )


async def confirmed_m05() -> OrderCase:
    """M05 (KMH, Japanese): every line the convert list did not map gets the
    candidate Sales confirms."""
    case = await in_review("M05")
    for line in case.lines:
        if not line.mapping.ready:
            case = await intake(OrderCases()).confirm_mapping(
                SCOPE, case, line.line_no, line.mapping.candidates[0], Actor(user_id=AN), T0
            )
    return confirmed(decided_all(case))


async def correction_requested_m04() -> OrderCase:
    case = await in_review("M04")
    case = case.dispose("price_mismatch:2", AskCustomer(by=AN, at=T0), Actor(user_id=AN))
    return case.request_correction()


async def change_review_m07() -> OrderCase:
    """M04 in Bravo, then its Rev.1 (M07): a change to apply in Bravo."""
    cases = OrderCases()
    case = await in_review("M04", cases)
    case = await intake(cases).confirm_mapping(SCOPE, case, 3, "CB-2008", Actor(user_id=AN), T0)
    cases.cases[case.case_id] = uploaded(decided_all(case))
    message = await MockInbox.load(SCOPE).get_message(SCOPE, "M07")
    assert message is not None
    outcome = await intake(cases).process(SCOPE, message)
    assert outcome.case is not None and outcome.case.status is OrderStatus.CHANGE_REVIEW
    return outcome.case


async def cannot_supply_m01() -> OrderCase:
    return (await in_review("M01")).close(CloseReason.CANNOT_SUPPLY, Actor(user_id=AN), T0)


# ------------------------------------------------------------ quote cases --


def _decision() -> PricingDecision:
    return PricingDecision(
        decided_by=DIEU,
        decided_at=T0,
        lme=LmeMonth(month="2026-09", usd_per_tonne=Decimal(10870)),
        lines=(
            LinePrice(
                line_no=1,
                unit_price=DECIDED,
                moq=Decimal(3000),
                lead_time_days=45,
                copper_basis=BAND,
            ),
        ),
        management_guidance=GUIDANCE,
    )


async def m10(step: str) -> QuoteCase:
    """M10 (KMH, Japanese) walked to ``step``."""
    cases = QuoteCases()
    quotation = quotation_service(cases)
    case = await quotation.open_case(SCOPE, "M10", uuid.uuid4())
    if step == "received":
        return case
    case = case.draft_ycbg()
    if step == "ycbg_drafted":
        return case
    case = case.record_ycbg(YCBG_M10, by=DIEU, at=T0)
    if step == "ycbg_recorded":
        return case
    case = case.send_to_design()
    cases.cases.append(case)
    reply = await quotation.take_design_reply(SCOPE, "M25")
    assert isinstance(reply, ReplyAttached)
    case = reply.case
    if step == "design_replied":
        return case
    if step == "declined":
        return case.decline_request(by=DIEU, at=T0, reason=DeclineReason.DESIGN_CANNOT)
    if step == "spec_discussion":
        return case.discuss_spec()
    case = await quotation.decide_price(SCOPE, case, _decision(), by=DIEU)
    case = await quotation.submit(
        SCOPE, case, quote_no=QUOTE_NO, issued_on=T0.date(), by=DIEU, at=T0
    )
    if step == "pending_approval":
        return case
    assert case.submission is not None
    case = case.approve(
        QuoteActor(user_id=GIANG, capabilities=frozenset({QuoteCapability.APPROVE})),
        at=NOW,
        document_sha256=case.submission.document_sha256,
    )
    if step == "approved":
        return case
    case = case.mark_sent(by=DIEU, at=NOW)
    assert step == "sent"
    return case


# Per kind: the case at the kind's state, and one before it.
type Builder = Callable[[], Awaitable[OrderCase | QuoteCase]]


async def _m01() -> OrderCase:
    return await in_review("M01")


async def _m04() -> OrderCase:
    return await in_review("M04")


async def _confirmed_m01() -> OrderCase:
    return confirmed(await in_review("M01"))


async def _confirmed_m13() -> OrderCase:
    return confirmed(decided_all(await in_review("M13")))


async def _uploaded_m27() -> OrderCase:
    return uploaded(decided_all(await typed_all(await in_review("M27"))))


async def _candidate_m05() -> OrderCase:
    case = await in_review("M05")
    (line,) = [line for line in case.lines if line.mapping.status is MappingStatus.AMBIGUOUS]
    return await intake(OrderCases()).confirm_mapping(
        SCOPE, case, line.line_no, line.mapping.candidates[0], Actor(user_id=AN), T0
    )


def _quote(step: str) -> Builder:
    async def build() -> QuoteCase:
        return await m10(step)

    return build


KINDS: dict[str, tuple[Builder, Builder]] = {
    "bravo_upload": (lambda: _prepared_m01(), _m01),
    "cross_check_sheet": (_uploaded_m27, _m01),
    "confirmation_draft": (_confirmed_m01, lambda: _cross_checked_m01()),
    "portal_checklist": (_confirmed_m13, lambda: _uploaded_m13()),
    "correction_request": (correction_requested_m04, _m04),
    "change_summary": (change_review_m07, _m04),
    "convert_list_proposal": (_candidate_m05, _m01),
    "design_code_request": (_m04, _m01),
    "cannot_supply_draft": (cannot_supply_m01, _m01),
    "ycbg_draft": (_quote("ycbg_drafted"), _quote("received")),
    "design_request_draft": (_quote("ycbg_recorded"), _quote("ycbg_drafted")),
    "spec_discussion_draft": (_quote("spec_discussion"), _quote("design_replied")),
    "decline_draft": (_quote("declined"), _quote("design_replied")),
    "quotation_preview": (_quote("pending_approval"), _quote("design_replied")),
    "quotation_final": (_quote("approved"), _quote("pending_approval")),
    "send_draft": (_quote("approved"), _quote("pending_approval")),
    "master_list_update": (_quote("sent"), _quote("approved")),
}


async def _prepared_m01() -> OrderCase:
    return prepared(await in_review("M01"))


async def _cross_checked_m01() -> OrderCase:
    return cross_checked(await in_review("M01"))


async def _uploaded_m13() -> OrderCase:
    return cross_checked(decided_all(await in_review("M13")))


def _kind_of(case: OrderCase | QuoteCase) -> CaseKind:
    return CaseKind.ORDER if isinstance(case, OrderCase) else CaseKind.QUOTE


async def render(
    world: World,
    case: OrderCase | QuoteCase,
    kind: str,
    *,
    as_: AccessContext | None = None,
    svc: ArtifactService | None = None,
) -> list[ArtifactView]:
    world.put(case)
    svc = svc or service(world)
    who = as_ or context(DIEU if kind == "quotation_preview" else AN)
    if isinstance(case, OrderCase):
        return await svc.render_order(who, case.case_id, kind=kind, case_version=case.case_version)
    return await svc.render_quote(who, case.case_id, kind=kind, case_version=case.case_version)


async def download(
    world: World, case: OrderCase | QuoteCase, artifact_id: uuid.UUID, *, as_: AccessContext
) -> bytes:
    svc = service(world)
    if isinstance(case, OrderCase):
        return (await svc.order_artifact(as_, case.case_id, artifact_id)).data
    return (await svc.quote_artifact(as_, case.case_id, artifact_id)).data


def _record(case: OrderCase | QuoteCase, kind: str) -> ArtifactRecord:
    return ArtifactRecord(
        artifact_id=uuid.uuid4(),
        case_kind=_kind_of(case),
        case_id=case.case_id,
        case_version=case.case_version,
        kind=kind,
        template_ref="sales_test@1.0.0",
        sha256=hashlib.sha256(b"file").hexdigest(),
        content_type="application/octet-stream",
        size_bytes=4,
        created_by=AN,
        created_at=T0,
    )


def _bumped[CaseT: OrderCase | QuoteCase](case: CaseT) -> CaseT:
    """The same case one edit later, still in its state: what any change to
    a value, a mapping or a disposition does to the version."""
    bumped = type(case).model_validate({**dict(case), "case_version": case.case_version + 1})
    return cast(CaseT, bumped)


def _who(kind: str) -> AccessContext:
    return context(DIEU if kind == "quotation_preview" else AN)


# --------------------------------------------------------------- the table --


def test_every_kind_in_the_gate_tables_is_tested_here() -> None:
    assert set(KINDS) == set(ORDER_GATES) | set(QUOTE_GATES)


@pytest.mark.parametrize("kind", sorted(KINDS))
async def test_each_kind_renders_at_its_state_with_template_version_and_hash(kind: str) -> None:
    world = World()
    case = await KINDS[kind][0]()

    views = await render(world, case, kind)

    assert views
    for view in views:
        record = world.artifacts[(SCOPE, view.artifact_id)]
        data = world.objects[record.object_key(SCOPE)]
        assert record.object_key(SCOPE) == (
            f"{SCOPE.tenant_id}/{SCOPE.workspace_id}/sales/{case.case_id}/{record.artifact_id}"
        )
        assert (record.kind, record.case_version) == (kind, case.case_version)
        assert record.sha256 == hashlib.sha256(data).hexdigest() == view.sha256
        assert record.size_bytes == len(data)
        pinned = {
            ARTIFACT_COPY.emails.version,
            ARTIFACT_COPY.documents.version,
            ARTIFACT_COPY.upload.policy_version,
        }
        assert record.template_ref.rsplit("@", 1)[1] in pinned, record.template_ref
        assert record.content_type in (XLSX, PDF, EML)
        served = await download(world, case, record.artifact_id, as_=_who(kind))
        assert served == data
    assert [e.action for e in world.audit].count("sales.artifact.rendered") == len(views)


@pytest.mark.parametrize("kind", sorted(KINDS))
async def test_each_kind_is_refused_before_its_state(kind: str) -> None:
    world = World()
    before = await KINDS[kind][1]()
    world.put(before)
    stored = _record(before, kind)
    world.artifacts[(SCOPE, stored.artifact_id)] = stored
    world.objects[stored.object_key(SCOPE)] = b"file"

    with pytest.raises(ConflictError) as rendering:
        await render(world, before, kind)
    with pytest.raises(ConflictError) as downloading:
        await download(world, before, stored.artifact_id, as_=_who(kind))

    assert rendering.value.details["kind"] == kind
    assert downloading.value.details["kind"] == kind
    assert "reason" not in downloading.value.details  # the gate, not the version


@pytest.mark.parametrize("kind", sorted(KINDS))
async def test_each_kind_is_refused_after_an_edit_bumped_the_case_version(kind: str) -> None:
    world = World()
    case = await KINDS[kind][0]()
    (view, *_) = await render(world, case, kind)
    world.put(_bumped(case))

    with pytest.raises(ConflictError) as refused:
        await download(world, case, view.artifact_id, as_=_who(kind))

    assert refused.value.details["reason"] == "stale_artifact"
    assert refused.value.details["artifact_case_version"] == case.case_version
    # Rendered again at the new version, it is served.
    (fresh, *_) = await render(world, _bumped(case), kind)
    assert fresh.case_version == case.case_version + 1
    assert await download(world, case, fresh.artifact_id, as_=_who(kind))


async def test_a_render_names_the_version_the_caller_saw() -> None:
    world = World()
    case = await _prepared_m01()
    world.put(_bumped(case))

    with pytest.raises(ConflictError, match="đã thay đổi"):
        await service(world).render_order(
            context(), case.case_id, kind="bravo_upload", case_version=case.case_version
        )


async def test_rendering_again_at_the_same_version_returns_what_is_stored() -> None:
    world = World()
    case = await _prepared_m01()

    first = await render(world, case, "bravo_upload")
    again = await render(world, case, "bravo_upload")

    assert [v.artifact_id for v in again] == [v.artifact_id for v in first]
    assert len(world.objects) == 1


async def test_two_renders_at_once_record_one_set_of_files() -> None:
    """Both find nothing stored and draw their files; the one that records
    second records nothing, removes the objects it stored, and answers the
    files of the one that recorded first (ticket 06)."""
    world = World()
    case = await m10("pending_approval")
    other: list[ArtifactView] = []

    async def another_render() -> None:
        other.extend(await render(world, case, "quotation_preview"))

    world.meanwhile = another_render
    answered = await render(world, case, "quotation_preview")

    assert {v.content_type for v in other} == {XLSX, PDF}
    assert answered == other
    assert {r.artifact_id for r in world.artifacts.values()} == {v.artifact_id for v in other}
    assert set(world.objects) == {
        world.artifacts[(SCOPE, v.artifact_id)].object_key(SCOPE) for v in other
    }
    assert [e.resource_id for e in world.audit] == [str(v.artifact_id) for v in other]


async def test_a_render_whose_commit_fails_leaves_no_object() -> None:
    world = World()
    case = await _prepared_m01()
    world.commit_fails = "before"

    with pytest.raises(InfrastructureError):
        await render(world, case, "bravo_upload")

    assert (world.artifacts, world.objects, world.audit) == ({}, {}, [])


async def test_a_commit_that_went_through_keeps_its_objects() -> None:
    """A commit can fail on the way back after it was made: the objects of
    records the store holds stay, or their downloads would answer 503."""
    world = World()
    case = await _prepared_m01()
    world.commit_fails = "after"

    with pytest.raises(InfrastructureError):
        await render(world, case, "bravo_upload")

    (record,) = world.artifacts.values()
    assert set(world.objects) == {record.object_key(SCOPE)}


async def test_the_eval_worlds_log_records_one_file_of_a_kind_per_version() -> None:
    """The eval world stands in for `sales.artifacts` and its unique key: a
    second record of one kind and content type at one version is not
    recorded; another content type, or the next version, is."""
    first = _record(await _prepared_m01(), "bravo_upload")
    again = first.model_copy(update={"artifact_id": uuid.uuid4()})
    pdf = again.model_copy(update={"artifact_id": uuid.uuid4(), "content_type": "application/pdf"})
    later = again.model_copy(
        update={"artifact_id": uuid.uuid4(), "case_version": first.case_version + 1}
    )
    async with MemoryStore()(SCOPE) as work:
        assert await work.artifacts.add(first)
        assert not await work.artifacts.add(again)
        assert await work.artifacts.get(again.artifact_id) is None
        assert await work.artifacts.add(pdf)
        assert await work.artifacts.add(later)


async def test_an_unknown_kind_is_refused() -> None:
    with pytest.raises(DomainError):
        await render(World(), await _prepared_m01(), "invoice")


async def test_a_bravo_upload_is_not_rendered_for_a_change_to_apply_in_bravo() -> None:
    """A revision of an order in Bravo is applied there; a new upload file
    would key the order a second time."""
    with pytest.raises(ConflictError):
        await render(World(), await change_review_m07(), "bravo_upload")


# ------------------------------------------------------ price and tenancy --


async def test_a_price_bearing_kind_needs_the_price_scope_to_render_or_download() -> None:
    world = World()
    case = await _prepared_m01()
    (view,) = await render(world, case, "bravo_upload")
    no_prices = context(AN, PIC_SCOPES - {"sales.price.read"})

    with pytest.raises(PermissionDeniedError) as rendering:
        await render(world, case, "bravo_upload", as_=no_prices)
    with pytest.raises(PermissionDeniedError) as downloading:
        await download(world, case, view.artifact_id, as_=no_prices)

    assert rendering.value.details["action"] == downloading.value.details["action"]
    assert downloading.value.details["action"] == "sales.price.read"


async def test_a_kind_without_prices_needs_no_price_scope() -> None:
    world = World()
    case = await _candidate_m05()
    no_prices = context(AN, PIC_SCOPES - {"sales.price.read"})

    (view,) = await render(world, case, "convert_list_proposal", as_=no_prices)

    assert await download(world, case, view.artifact_id, as_=no_prices)


async def test_another_tenants_case_and_artifact_are_not_found() -> None:
    world = World()
    case = await _prepared_m01()
    (view,) = await render(world, case, "bravo_upload")
    beta = context(AN, scope=OTHER)

    with pytest.raises(NotFoundError):
        await download(world, case, view.artifact_id, as_=beta)
    with pytest.raises(NotFoundError):
        await service(world).render_order(
            beta, case.case_id, kind="bravo_upload", case_version=case.case_version
        )


async def test_another_cases_artifact_is_not_found() -> None:
    world = World()
    first = await _prepared_m01()
    second = await _prepared_m01()
    world.put(second)
    (view,) = await render(world, first, "bravo_upload")

    with pytest.raises(NotFoundError):
        await download(world, second, view.artifact_id, as_=context())


async def test_bytes_that_are_not_the_recorded_ones_are_never_served() -> None:
    world = World()
    case = await _prepared_m01()
    (view,) = await render(world, case, "bravo_upload")
    record = world.artifacts[(SCOPE, view.artifact_id)]
    world.objects[record.object_key(SCOPE)] += b"tampered"

    with pytest.raises(InfrastructureError):
        await download(world, case, view.artifact_id, as_=context())


async def test_render_and_download_are_audited_with_ids_and_hashes_only() -> None:
    world = World()
    case = await _prepared_m01()
    (view,) = await render(world, case, "bravo_upload")
    await download(world, case, view.artifact_id, as_=context())

    rendered, downloaded = world.audit
    assert (rendered.action, downloaded.action) == (
        "sales.artifact.rendered",
        "sales.artifact.downloaded",
    )
    assert set(downloaded.details) == {
        "case_kind",
        "case_id",
        "case_version",
        "kind",
        "template_ref",
        "sha256",
    }
    text = repr([e.details for e in world.audit])
    for line in case.lines:
        assert str(line.po_line.unit_price) not in text


# ----------------------------------------------------------------- contents --


def _sheet_rows(data: bytes, sheet: int = 0) -> list[list[Any]]:
    book = load_workbook(io.BytesIO(data))
    return [[c.value for c in row] for row in book.worksheets[sheet].iter_rows()]


def _xlsx_text(data: bytes) -> str:
    book = load_workbook(io.BytesIO(data))
    return "\n".join(
        str(cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def _eml(data: bytes) -> email.message.EmailMessage:
    message = email.message_from_bytes(data, policy=policy.default)
    assert isinstance(message, email.message.EmailMessage)
    return message


def _body(data: bytes) -> str:
    part = _eml(data).get_body(("plain",))
    assert part is not None
    return str(part.get_content())


async def _rendered(case: OrderCase | QuoteCase, kind: str) -> list[tuple[ArtifactRecord, bytes]]:
    world = World()
    views = await render(world, case, kind)
    records = [world.artifacts[(SCOPE, v.artifact_id)] for v in views]
    return [(r, world.objects[r.object_key(SCOPE)]) for r in records]


async def _one(case: OrderCase | QuoteCase, kind: str) -> bytes:
    ((_, data),) = await _rendered(case, kind)
    return data


async def test_the_upload_file_holds_the_decided_lines_in_the_mock_layout() -> None:
    case = await _prepared_m01()
    data = await _one(case, "bravo_upload")

    header, *rows = _sheet_rows(data)
    layout = ARTIFACT_COPY.upload
    assert header == [c.header for c in layout.columns]
    fields = [c.field for c in layout.columns]
    assert len(rows) == len(case.lines)
    for row, line in zip(rows, case.lines, strict=True):
        values = dict(zip(fields, row, strict=True))
        assert values["prv_code"] == line.mapping.prv_code
        assert values["po_no"] == case.header.po_no
        assert Decimal(str(values["quantity"])) == line.po_line.quantity
        assert Decimal(str(values["unit_price"])) == line.po_line.unit_price
    assert "MOCK" in _xlsx_text(data)


async def test_a_description_starting_with_an_equals_sign_is_written_as_text() -> None:
    """M20 line 3's description is ``=HYPERLINK(...)``: a cell of type text
    with the quote prefix, never a formula (G29)."""
    case = await _prepared_m01_like("M20")
    data = await _one(case, "bravo_upload")

    sheet = load_workbook(io.BytesIO(data)).worksheets[0]
    column = [c.field for c in ARTIFACT_COPY.upload.columns].index("description") + 1
    cell = sheet.cell(row=1 + 3, column=column)
    assert str(cell.value).startswith("=HYPERLINK(")
    assert cell.data_type == "s"
    assert cell.quotePrefix is True
    raw = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in raw


async def _prepared_m01_like(message_id: str) -> OrderCase:
    return prepared(decided_all(await in_review(message_id)))


async def test_the_cross_check_sheet_names_who_typed_and_who_recorded_the_entry() -> None:
    case = await _uploaded_m27()
    data = await _one(case, "cross_check_sheet")
    text = _xlsx_text(data)

    assert "SO26-1001" in text
    assert "Nguyễn Văn An" in text  # recorded the Bravo entry and typed values
    rows = _sheet_rows(data)
    header_at = next(i for i, row in enumerate(rows) if row[:2] == ["Dòng", "Giá trị"])
    assert rows[header_at][2:4] == ["Trên PO", "Sales sửa thành"]
    body = rows[header_at + 1 :]
    typed = {row[0] for row in body if row[8] == "Nguyễn Văn An"}
    assert typed == {2, 4, 5, 6}  # the lines whose flagged values Sales typed
    prices = {row[0]: row for row in body if row[1] == "Đơn giá"}
    # Per value: as read, as Sales typed it, and in the upload file.
    read, corrected, upload = (prices[2][i] for i in (2, 3, 5))
    assert Decimal(str(read)) == case.line(2).po_line.unit_price
    assert Decimal(str(corrected)) == Decimal(str(upload)) == CORRECTED_PRICE
    for line in case.lines:
        if line.line_no == 2:
            continue
        assert Decimal(str(prices[line.line_no][5])) == line.po_line.unit_price
        # Typed as read on the flagged lines, nothing typed on the others.
        assert (prices[line.line_no][3] is None) == (line.line_no not in {4, 5, 6})


async def test_the_cross_check_sheet_lists_each_schedule_beside_its_upload_row() -> None:
    """Ticket 21: a line the PO splits over several dates has a row per
    schedule, as printed, where, and as the upload file's row holds it."""
    case = uploaded(await in_review("M34"))
    rows = _sheet_rows(await _one(case, "cross_check_sheet"))
    header_at = next(i for i, row in enumerate(rows) if row[:2] == ["Dòng", "Giá trị"])

    schedules = [row for row in rows[header_at + 1 :] if str(row[1]).startswith("Đợt giao")]

    assert [(row[0], row[1], row[2], row[4], row[5]) for row in schedules] == [
        (1, "Đợt giao 1/2", "2026-11-16: 6100", "M34-A1 PO!H11", "2026-11-16: 6100"),
        (1, "Đợt giao 2/2", "2026-12-01: 6100", "M34-A1 PO!H12", "2026-12-01: 6100"),
        (3, "Đợt giao 1/3", "2026-11-27: 1500", "M34-A1 PO!H15", "2026-11-27: 1500"),
        (3, "Đợt giao 2/3", "2026-12-11: 900", "M34-A1 PO!H16", "2026-12-11: 900"),
        (3, "Đợt giao 3/3", "2026-12-28: 600", "M34-A1 PO!H17", "2026-12-28: 600"),
    ]


async def test_the_upload_file_prints_the_value_sales_corrected_not_the_value_read() -> None:
    """Ticket 15: a corrected value reaches the upload file; the value as read
    stays on the line as the evidence."""
    case = await _uploaded_m27()
    data = await _one(case, "bravo_upload")

    _header, *rows = _sheet_rows(data)
    fields = [c.field for c in ARTIFACT_COPY.upload.columns]
    by_line = {int(row[fields.index("line_no")]): row for row in rows}
    price = fields.index("unit_price")
    assert Decimal(str(by_line[2][price])) == CORRECTED_PRICE
    assert case.line(2).po_line.unit_price != CORRECTED_PRICE
    assert Decimal(str(by_line[1][price])) == case.line(1).po_line.unit_price


async def test_the_confirmation_names_the_pic_the_checker_and_a_date_per_line() -> None:
    case = await _confirmed_m01()
    message = _eml(await _one(case, "confirmation_draft"))
    body = _body(await _one(case, "confirmation_draft"))

    assert "Nguyễn Văn An" in body and "Hoàng Thị Diệu" in body
    assert body.count(CONFIRMED_ON.strftime("%d/%m/%Y")) == len(case.lines)
    assert message["X-Unsent"] == "1"
    assert message["From"].addresses[0].addr_spec == "an.nguyen@alpha.local"
    assert message["Reply-To"].addresses[0].addr_spec == "sales@seller.example"
    customer = (await MockSalesCatalog.load(SCOPE).customer_by_code(SCOPE, "VLX")).data
    assert customer is not None
    assert {a.addr_spec for a in message["To"].addresses} == {c.address for c in customer.contacts}


# M34 (VLX, Vietnamese): line 1 is split in two, line 3 in three; line 2 has one date.
_M34_DATES = {
    1: [date(2026, 11, 18), date(2026, 12, 2)],
    2: [date(2026, 11, 27)],
    3: [date(2026, 11, 27), date(2026, 12, 11), date(2026, 12, 29)],
}


async def _confirmed_m34() -> OrderCase:
    checked = cross_checked(decided_all(await in_review("M34")))
    return checked.confirm(Actor(user_id=AN), T0, _M34_DATES)


async def test_a_split_line_is_confirmed_to_the_customer_schedule_by_schedule() -> None:
    """Ticket 35: each schedule's quantity with the date confirmed for it, not
    the line's whole quantity on one day."""
    case = await _confirmed_m34()
    body = _body(await _one(case, "confirmation_draft"))
    lines = {
        int(row.split(":")[0].removeprefix("- Dòng ")): row
        for row in body.splitlines()
        if row.startswith("- Dòng")
    }

    assert lines[1].endswith(
        "ngày giao xác nhận theo 2 đợt: 6100 m ngày 18/11/2026; 6100 m ngày 02/12/2026"
    )
    assert lines[3].endswith(
        "ngày giao xác nhận theo 3 đợt: 1500 m ngày 27/11/2026; 900 m ngày 11/12/2026;"
        " 600 m ngày 29/12/2026"
    )
    assert lines[2].endswith("ngày giao xác nhận 27/11/2026")
    assert "đợt" not in lines[2]


async def test_a_portal_checklist_has_a_row_per_schedule_of_a_split_line() -> None:
    """VLX confirms by email; the checklist is drawn for it as if on a portal."""
    case = await _confirmed_m34()
    customer = (await MockSalesCatalog.load(SCOPE).customer_by_code(SCOPE, "VLX")).data
    assert customer is not None
    drafting = Drafting(
        case=case,
        case_version=case.case_version,
        customer=customer.model_copy(update={"confirmation_channel": "portal"}),
        people=People({m.user_id: m for m in MEMBERS}),
        drafter=AN,
        at=NOW,
        copy=ARTIFACT_COPY,
        writer=WRITER,
        scope=SCOPE,
    )

    (rendered,) = portal_checklist(drafting)

    rows = _sheet_rows(rendered.data)
    header_at = next(i for i, row in enumerate(rows) if row and row[0] == "Dòng")
    listed = [(row[0], Decimal(str(row[3])), row[7]) for row in rows[header_at + 1 :] if row[0]]
    as_dates = [
        (no, qty, day.date() if isinstance(day, datetime) else day) for no, qty, day in listed
    ]
    assert as_dates == [
        (1, Decimal(6100), date(2026, 11, 18)),
        (1, Decimal(6100), date(2026, 12, 2)),
        (2, case.line(2).effective.quantity, date(2026, 11, 27)),
        (3, Decimal(1500), date(2026, 11, 27)),
        (3, Decimal(900), date(2026, 12, 11)),
        (3, Decimal(600), date(2026, 12, 29)),
    ]


async def test_m35s_short_split_line_is_confirmed_per_schedule_once_pc_agreed() -> None:
    """M35 (BRN, PDF): line 2's schedules are too soon. PC's agreement is per
    line, covering every schedule its finding names; then each schedule is
    confirmed on its own date."""
    review = await in_review("M35")
    # Line 1's schedules add up to 5,490 where it prints 6,100: Sales takes the
    # schedules' sum, as the customer's PO meant.
    review = review.dispose(
        "line_total_mismatch:-",
        CorrectedBySales(value="5490", source="PO gốc", by=AN, at=T0),
        Actor(user_id=AN),
    )
    checked = cross_checked(decided_all(review))
    short = next(f for f in checked.findings if f.code is FindingCode.REQUESTED_DATE_SHORT_LT)
    assert short.line_no == 2
    dates = {
        line.line_no: [CONFIRMED_ON + timedelta(days=7 * n) for n in range(line.deliveries)]
        for line in checked.lines
    }
    assert len(dates[2]) == checked.line(2).deliveries > 1

    with pytest.raises(ConflictError, match="no PC-confirmed date"):
        checked.confirm(Actor(user_id=AN), T0, dates)
    agreed = checked.record_pc_confirmation(2, Actor(user_id=AN), T0)
    case = agreed.confirm(Actor(user_id=AN), T0, dates)

    body = _body(await _one(case, "confirmation_draft"))
    (line_2,) = [row for row in body.splitlines() if row.startswith("- Dòng 2:")]
    schedules = case.line(2).effective.schedules
    assert line_2.endswith(
        f"theo {len(schedules)} đợt: "
        + "; ".join(
            f"{s.quantity} {case.line(2).effective.uom} ngày {day:%d/%m/%Y}"
            for s, day in zip(schedules, dates[2], strict=True)
        )
    )


async def test_a_customer_who_confirms_on_their_portal_gets_a_checklist_not_an_email() -> None:
    case = await _confirmed_m13()  # NRV confirms on its portal

    with pytest.raises(ConflictError) as refused:
        await _one(case, "confirmation_draft")
    data = await _one(case, "portal_checklist")

    assert refused.value.details["reason"] == "confirmation_channel"
    assert "Đã nhập trên portal" in _xlsx_text(data)
    listed = World()
    listed.put(case)
    available = (await service(listed).order_artifacts(context(), case.case_id)).available
    assert "portal_checklist" in available and "confirmation_draft" not in available


async def test_the_correction_request_lists_exactly_the_findings_sent_back() -> None:
    case = await correction_requested_m04()
    body = _body(await _one(case, "correction_request"))

    asked = [f for f in case.findings if f.disposition.kind.value == "ask_customer"]
    assert [f.key for f in asked] == ["price_mismatch:2"]
    assert "Line 2" in body and "0.658" in body and "0.6980" in body  # NRV writes English
    assert "NV-CB7-20-OR" not in body  # line 3's finding was not sent back
    assert body.count("\n- ") == 1


async def test_the_correction_request_asks_for_a_po_typed_in_a_legacy_font_in_unicode() -> None:
    """Ticket 20: M33 (VLX writes Vietnamese) is typed in `.VnTime`; the draft
    names the whole PO and the first cell typed in that font."""
    case = await in_review("M33")
    asked = case.dispose(
        "legacy_encoding:-", AskCustomer(by=AN, at=T0), Actor(user_id=AN)
    ).request_correction()
    body = _body(await _one(asked, "correction_request"))

    assert "Toàn bộ PO" in body and "font tiếng Việt cũ (TCVN3/VNI)" in body
    assert "Trên PO: TCVN3; giá trị đúng: Unicode" in body
    assert "PO!C10" in body


async def test_the_change_summary_lists_what_the_revision_changed() -> None:
    case = await change_review_m07()
    rows = _sheet_rows(await _one(case, "change_summary"))

    assert case.changes
    listed = {(row[0], row[2], row[3]) for row in rows if isinstance(row[0], int)}
    assert listed == {(c.line_no, c.before, c.after) for c in case.changes}


async def test_the_design_code_note_lists_the_line_without_a_code() -> None:
    body = _body(await _one(await _m04(), "design_code_request"))

    assert "NV-CB7-20-OR" in body
    assert "NV-HW22-BK" not in body  # mapped by the convert list


async def test_a_missing_language_refuses_to_render_and_stores_nothing() -> None:
    """KMH writes Japanese: with no ``ja`` send draft the copy refuses, with
    the reason, rather than sending Vietnamese (G33)."""
    emails = ARTIFACT_COPY.emails
    without_ja = SalesEmailCopy.model_validate(
        {
            **emails.model_dump(by_alias=True),
            "emails": {
                kind: {
                    lang: t
                    for lang, t in by_language.items()
                    if kind != "send_draft" or lang != "ja"
                }
                for kind, by_language in emails.model_dump(by_alias=True)["emails"].items()
            },
        }
    )
    copy = platform_book(
        QUOTE_RULES, PRICING, without_ja, ARTIFACT_COPY.documents, ARTIFACT_COPY.upload
    )
    world = World()
    case = await m10("approved")

    with pytest.raises(ConflictError) as refused:
        await render(world, case, "send_draft", svc=service(world, copy))

    assert refused.value.details == {
        "kind": "send_draft",
        "language": "ja",
        "copy": "sales_emails@1.2.0",
        "reason": "template_missing",
    }
    assert not world.objects and not world.artifacts


@pytest.mark.parametrize(
    ("build", "language", "expected"),
    [
        (_confirmed_m01, "vi", "Kính gửi"),
        (lambda: _confirmed_cvg(), "en", "Dear"),
        (confirmed_m05, "ja", "御中"),
    ],
)
async def test_an_order_confirmation_renders_in_each_language(
    build: Builder, language: str, expected: str
) -> None:
    body = _body(await _one(await build(), "confirmation_draft"))

    assert expected in body


async def _confirmed_cvg() -> OrderCase:
    return confirmed(decided_all(await in_review("M02")))


# ---------------------------------------------------------------- quotation --


async def test_m10_quotation_is_japanese_and_from_the_submitted_document_only() -> None:
    case = await m10("approved")
    files = {r.content_type: data for r, data in await _rendered(case, "quotation_final")}

    pdf_text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(files[PDF])).pages)
    xlsx_text = _xlsx_text(files[XLSX])
    assert case.submission is not None and case.approval is not None
    for text in (pdf_text, xlsx_text):
        assert "御見積書" in text and QUOTE_NO in text
        assert "0.689" in text  # 0.6890: a number in the xlsx, printed in full in the PDF
        assert case.submission.document_sha256 in text
        assert "Đỗ Minh Giang" in text  # the approver, by name
    assert {r.template_ref for r, _ in await _rendered(case, "quotation_final")} == {
        "sales_documents.quotation.ja@1.2.0"
    }


async def test_the_preview_says_it_is_not_approved_and_is_only_for_pricer_and_approvers() -> None:
    case = await m10("pending_approval")
    files = {r.content_type: d for r, d in await _rendered(case, "quotation_preview")}
    assert "プレビュー" in _xlsx_text(files[XLSX])

    world = World()
    with pytest.raises(ConflictError):  # An neither priced it nor approves
        await render(world, case, "quotation_preview", as_=context(AN))
    assert await render(world, case, "quotation_preview", as_=context(DIEU))  # the pricer


async def test_the_send_draft_attaches_the_stored_final_files_and_names_the_spec() -> None:
    world = World()
    case = await m10("approved")
    finals = await render(world, case, "quotation_final")
    (draft,) = await render(world, case, "send_draft")

    message = _eml(world.objects[world.artifacts[(SCOPE, draft.artifact_id)].object_key(SCOPE)])
    attached = {
        hashlib.sha256(part.get_content()).hexdigest() for part in message.iter_attachments()
    }
    assert attached == {view.sha256 for view in finals}
    assert case.submission is not None
    (line,) = case.submission.document.lines
    assert line.spec_no in _body(
        world.objects[world.artifacts[(SCOPE, draft.artifact_id)].object_key(SCOPE)]
    )
    assert {a.addr_spec for a in message["To"].addresses} == {
        r.address for r in case.submission.document.recipients
    }


async def test_no_customer_facing_artifact_carries_another_customers_price_or_guidance() -> None:
    """M10's evidence holds other customers' prices and the reference, and its
    decision a management instruction: none reaches the customer."""
    texts: list[str] = []
    for step, kinds in (
        ("approved", ("quotation_final", "send_draft")),
        ("pending_approval", ("quotation_preview",)),
        ("spec_discussion", ("spec_discussion_draft",)),
        ("declined", ("decline_draft",)),
    ):
        case = await m10(step)
        for kind in kinds:
            for record, data in await _rendered(case, kind):
                if record.content_type == XLSX:
                    texts.append(_xlsx_text(data))
                elif record.content_type == PDF:
                    texts.append(
                        "".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
                    )
                else:
                    message = _eml(data)
                    texts.append(_body(data))
                    for part in message.iter_attachments():
                        content = part.get_content()
                        texts.append(
                            _xlsx_text(content)
                            if str(part.get_filename()).endswith(".xlsx")
                            else "".join(
                                p.extract_text() for p in PdfReader(io.BytesIO(content)).pages
                            )
                        )
    joined = "\n".join(texts)
    assert "0.689" in joined  # the check can see a price when there is one
    for internal in INTERNAL_PRICES:
        # As printed, and as a spreadsheet number reads back (0.7120 -> 0.712).
        for form in (internal, str(Decimal(internal).normalize())):
            assert form not in joined, form
    assert GUIDANCE not in joined


async def test_the_master_list_update_holds_the_sent_quotation_rows() -> None:
    case = await m10("sent")
    text = _xlsx_text(await _one(case, "master_list_update"))

    rows: tuple[Quotation, ...] = case.master_list_rows()
    for row in rows:
        assert row.quote_no in text and row.prv_code in text
    assert "MOCK" in text


async def test_the_decline_draft_gives_the_reason_in_the_customers_language() -> None:
    body = _body(await _one(await m10("declined"), "decline_draft"))

    assert "弊社設計が対応できない" in body


async def test_the_ycbg_and_the_design_request_go_to_design_in_vietnamese() -> None:
    form = _xlsx_text(await _one(await m10("ycbg_recorded"), "ycbg_draft"))
    message = _eml(await _one(await m10("ycbg_recorded"), "design_request_draft"))

    assert YCBG_M10 in form and "MOCK" in form
    assert [a.addr_spec for a in message["To"].addresses] == list(ARTIFACT_COPY.design_mailboxes)
    assert YCBG_M10 in str(message["Subject"])


# ------------------------------------- the tenant's own copy (ticket 36) --


def _tenant_book() -> TenantRuleBook:
    """The platform's copy and quote rules, and room for a tenant's own."""
    return platform_book(
        QUOTE_RULES, PRICING, ARTIFACT_COPY.emails, ARTIFACT_COPY.documents, ARTIFACT_COPY.upload
    )


def _seed(book: TenantRuleBook, scope: SalesScope, policy: Policy) -> None:
    assert isinstance(book.stored, MemoryRuleVersions)
    book.stored.seed_approved(scope.tenant_id.value, policy)


def _documents(version: str, title: str) -> SalesDocumentCopy:
    copy = ARTIFACT_COPY.documents.model_dump(mode="json", by_alias=True)
    for template in copy["documents"]["quotation"].values():
        template["title"] = title
    return SalesDocumentCopy.model_validate({**copy, "version": version})


async def test_drafts_are_written_in_the_tenants_own_copy_and_record_it() -> None:
    book = _tenant_book()
    emails = ARTIFACT_COPY.emails.model_dump(mode="json", by_alias=True)
    emails["emails"]["confirmation_draft"]["vi"]["subject"] = "PO ${po_no} đã xác nhận (demo)"
    _seed(book, SCOPE, SalesEmailCopy.model_validate({**emails, "version": "9.0.1"}))
    world = World()
    case = await _confirmed_m01()

    (view,) = await render(world, case, "confirmation_draft", svc=service(world, book))

    data = world.objects[world.artifacts[(SCOPE, view.artifact_id)].object_key(SCOPE)]
    assert view.template_ref == "sales_emails.confirmation_draft.vi@9.0.1"
    assert _eml(data)["Subject"] == f"PO {case.header.po_no} đã xác nhận (demo)"
    other = await book.artifact_copy(OTHER, [])
    assert other.emails.ref == ARTIFACT_COPY.emails.ref


async def test_the_send_draft_attaches_the_final_files_drawn_before_a_new_documents_copy() -> None:
    """A documents version approved between the final quotation and its send
    draft does not reach the attachments: files drawn at one case version
    read the copy the first of them stamped."""
    book = _tenant_book()
    world = World()
    svc = service(world, book)
    case = await m10("approved")
    finals = await render(world, case, "quotation_final", svc=svc)
    _seed(book, SCOPE, _documents("9.0.1", "BÁO GIÁ MỚI"))

    (draft,) = await render(world, case, "send_draft", svc=svc)

    message = _eml(world.objects[world.artifacts[(SCOPE, draft.artifact_id)].object_key(SCOPE)])
    attached = {
        hashlib.sha256(part.get_content()).hexdigest() for part in message.iter_attachments()
    }
    assert attached == {view.sha256 for view in finals}
    assert {view.template_ref for view in finals} == {"sales_documents.quotation.ja@1.2.0"}
    # The next case version draws with the tenant's version.
    assert (await book.artifact_copy(SCOPE, [])).documents.ref == "sales_documents@9.0.1"


async def test_a_file_whose_stamped_copy_is_no_longer_held_is_refused() -> None:
    book = _tenant_book()
    with pytest.raises(ConflictError) as refused:
        await book.artifact_copy(SCOPE, ["sales_documents.quotation.ja@0.9.0"])
    assert refused.value.details == {
        "copy": "sales_documents@0.9.0",
        "reason": "template_version_gone",
    }


async def test_the_design_drafts_go_to_the_tenants_current_design_mailboxes() -> None:
    book = _tenant_book()
    moved = QUOTE_RULES.model_copy(
        update={"policy_version": "9.0.1", "design_mailboxes": ("thietke@seller.example",)}
    )
    _seed(book, SCOPE, QuoteRules.model_validate(moved.model_dump()))
    world = World()

    (view,) = await render(
        world, await m10("ycbg_recorded"), "design_request_draft", svc=service(world, book)
    )

    message = _eml(world.objects[world.artifacts[(SCOPE, view.artifact_id)].object_key(SCOPE)])
    assert [a.addr_spec for a in message["To"].addresses] == ["thietke@seller.example"]


# ----------------------------------------------------------------- the copy --


def test_every_customer_draft_has_every_language_and_every_internal_one_vietnamese() -> None:
    customer = {
        "confirmation_draft",
        "correction_request",
        "cannot_supply_draft",
        "spec_discussion_draft",
        "decline_draft",
        "send_draft",
    }
    emails = ARTIFACT_COPY.emails.emails
    for kind in customer:
        assert set(emails[kind]) == {"vi", "en", "ja"}, kind
    for kind in set(emails) - customer:
        assert set(emails[kind]) == {"vi"}, kind
    assert set(ARTIFACT_COPY.documents.documents["quotation"]) == {"vi", "en", "ja"}


def test_a_template_naming_a_value_the_draft_lacks_refuses_to_render() -> None:
    from dw_sales.application.artifact_content import fill

    with pytest.raises(ConflictError) as refused:
        fill("Ngày ${confirmed_date}", {"po_no": "X"}, ref="t@1.0.0")

    assert refused.value.details["reason"] == "template_value_unknown"


def test_a_value_is_substituted_as_text_never_read_as_a_template() -> None:
    from dw_sales.application.artifact_content import fill

    assert fill("${a}", {"a": "${b} $c"}, ref="t@1.0.0") == "${b} $c"


def test_a_subject_with_a_line_break_is_refused_before_it_can_be_a_tenants_copy() -> None:
    """Ticket 36: a tenant's e-mail copy is proposed by a person; a subject is
    one header line, which the mail writer would refuse to build."""
    copy = ARTIFACT_COPY.emails.model_dump(mode="json", by_alias=True)
    copy["emails"]["send_draft"]["vi"]["subject"] = "Báo giá ${quote_no}\nBcc: x@example.com"

    with pytest.raises(ValidationError, match="subject"):
        SalesEmailCopy.model_validate(copy)


async def test_the_order_view_shows_each_schedules_confirmed_date_and_no_line_date() -> None:
    """Ticket 35: a split line's dates are on its schedules; a line delivered
    once keeps its one date on the line."""
    case = await _confirmed_m34()
    prices = PriceView(amounts=True, other_customers=False)

    split, single = order_line(case, case.line(1), prices), order_line(case, case.line(2), prices)

    assert [s.confirmed_delivery_date for s in split.schedules] == _M34_DATES[1]
    assert split.confirmed_delivery_date is None
    assert single.confirmed_delivery_date == _M34_DATES[2][0]
    assert [s.confirmed_delivery_date for s in single.schedules] == [None]
