"""`SourceViewPort`: one PDF page, or one sheet's grid, opened for a person.

Parsing only, pure Python, bounded by the caps reading uses (the order rules'
`IntakeCaps`, spec decision 13). A PDF page is copied out as a one-page PDF
for the browser to render (pdf.js): nothing is rasterised here. A sheet is
its cells with every flag a person could miss a value by: a hidden sheet,
row or column, a font the colour of its fill, a formula with no cached
value, and the colours themselves.
"""

from __future__ import annotations

import asyncio
import io
from datetime import date, datetime
from typing import Final, Literal

from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader, PdfWriter

from dw_sales.adapters.readers.documents import sniff
from dw_sales.adapters.readers.excel import font_matches_fill, open_for_viewing
from dw_sales.application.source import SheetCell, SheetGrid
from dw_sales.domain.order_checks import IntakeCaps

# What a grid shows of a sheet, whatever dimensions the sheet declares.
_MAX_ROWS: Final = 1000
_MAX_COLUMNS: Final = 52


class FileSourceView:
    """Implements `SourceViewPort`."""

    def __init__(self, caps: IntakeCaps) -> None:
        self._caps = caps

    def kind(self, data: bytes) -> Literal["pdf", "xlsx"] | None:
        match sniff(data):
            case "pdf":
                return "pdf"
            case "xlsx":
                return "xlsx"
            case _:
                return None

    async def pdf_page(self, data: bytes, page: int) -> tuple[bytes, int] | None:
        return await asyncio.to_thread(self._pdf_page, data, page)

    async def sheet_names(self, data: bytes) -> list[str]:
        values, _ = await asyncio.to_thread(open_for_viewing, data, self._caps)
        return [sheet.title for sheet in values.worksheets]

    async def sheet_grid(self, data: bytes, sheet: str) -> SheetGrid | None:
        return await asyncio.to_thread(self._grid, data, sheet)

    def _pdf_page(self, data: bytes, page: int) -> tuple[bytes, int] | None:
        if len(data) > self._caps.max_attachment_bytes:
            raise ValueError("the file is over the size read here")
        reader = PdfReader(io.BytesIO(data))
        count = len(reader.pages)
        if count > self._caps.max_pages:
            raise ValueError("the file has more pages than are read here")
        if not 1 <= page <= count:
            return None
        writer = PdfWriter()
        writer.add_page(reader.pages[page - 1])
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue(), count

    def _grid(self, data: bytes, title: str) -> SheetGrid | None:
        values, formulas = open_for_viewing(data, self._caps)
        if title not in values.sheetnames:
            return None
        sheet = values[title]
        formula_sheet = formulas[title]
        rows = min(sheet.max_row, _MAX_ROWS)
        columns = min(sheet.max_column, _MAX_COLUMNS)
        hidden_rows = {
            number for number, dimension in sheet.row_dimensions.items() if dimension.hidden
        }
        hidden_columns = _hidden_columns(sheet)
        cells: list[SheetCell] = []
        for row in sheet.iter_rows(min_row=1, max_row=rows, max_col=columns):
            for cell in row:
                if not isinstance(cell, Cell):
                    continue
                formula = formula_sheet[cell.coordinate]
                no_cache = (
                    isinstance(formula, Cell) and formula.data_type == "f" and cell.value is None
                )
                if cell.value is None and not no_cache:
                    continue
                cells.append(
                    SheetCell(
                        ref=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        text=_text(cell.value),
                        hidden_row=cell.row in hidden_rows,
                        hidden_column=cell.column in hidden_columns,
                        font_matches_fill=font_matches_fill(cell),
                        formula_without_cached_value=no_cache,
                        font_rgb=_rgb(
                            cell.font.color.rgb if cell.font and cell.font.color else None
                        ),
                        fill_rgb=_rgb(
                            cell.fill.fgColor.rgb
                            if cell.fill is not None and cell.fill.fill_type == "solid"
                            else None
                        ),
                    )
                )
        return SheetGrid(
            sheet=title,
            hidden_sheet=sheet.sheet_state != "visible",
            rows=rows,
            columns=columns,
            truncated=sheet.max_row > _MAX_ROWS or sheet.max_column > _MAX_COLUMNS,
            cells=cells,
        )


def _hidden_columns(sheet: Worksheet) -> set[int]:
    hidden: set[int] = set()
    for dimension in sheet.column_dimensions.values():
        if dimension.hidden and dimension.min is not None and dimension.max is not None:
            hidden.update(range(dimension.min, min(dimension.max, _MAX_COLUMNS) + 1))
    return hidden


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _rgb(value: object) -> str | None:
    return value if isinstance(value, str) and len(value) <= 8 else None
