"""Regenerate the mock email attachments from the fixtures in `data/`.

    uv run python -m dw_sales.adapters.mock.generate_attachments

The files stand in for what customers' own systems print, so the layouts and
the wording of item descriptions are theirs (documented in `README.md`), not
this context's. Purchase orders come from `data/purchase_orders.json`,
requests for quotation from `data/quote_requests.json` (the fixed-label
form) and `data/free_requests.json` (a customer's own PDF or spreadsheet,
which no reader reads and a model may, ticket 29), and Design's replies
from `data/design_replies.json`. A line that states no attributes is described
the way the catalogue describes the item its code maps to, so a description
never disagrees with the item it stands for, unless the fixture gives the
printed words verbatim.

Byte-for-byte deterministic: fixed document dates, fixed zip entry times and
no compression anywhere. A regeneration changes nothing unless a fixture did,
and a unit test proves the committed files are exactly what this produces.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Literal, NamedTuple, Self

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml import LXML
from pydantic import BaseModel, ConfigDict, Field, model_validator
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from dw_sales.adapters.mock.fixtures import (
    ATTACHMENTS_DIR,
    DATA_DIR,
    MOCK_ROOT,
    attachment_file_name,
    read_records,
)
from dw_sales.domain.catalog import (
    Colour,
    Conductor,
    ConvertEntry,
    Currency,
    Customer,
    Item,
    ItemFamily,
    Language,
    Packaging,
    Shield,
    Stranding,
)

_FROZEN = ConfigDict(frozen=True, extra="forbid")

# The fictional seller every mock document is addressed to.
SELLER_NAME = "Demo Wire & Cable Vietnam Co., Ltd."

# DejaVu Sans cut down to Latin and Vietnamese; licence and provenance in
# `fonts/`. Embedded (subset) into every PDF, so the text layer carries real
# Vietnamese characters rather than whatever a viewer substitutes.
FONT_PATH = MOCK_ROOT / "fonts" / "DejaVuSans-LatinVi.ttf"
_FONT = "DejaVuSans-LatinVi"
# The TCVN3 (ABC) font a customer still typing Vietnamese the pre-Unicode way
# sets on a cell: its letters are stored as other characters (ticket 20).
LEGACY_FONT = ".VnTime"

_FAMILY_WORDS: dict[ItemFamily, str] = {
    "hook_up_wire": "HOOK-UP WIRE",
    "multi_core_cable": "MULTI-CORE CABLE",
}
_CONDUCTOR_WORDS: dict[Conductor, str] = {"bare_copper": "BARE CU", "tinned_copper": "TINNED CU"}
_SHIELD_WORDS: dict[Shield, str] = {
    "none": "NO SHIELD",
    "foil": "FOIL SHIELD",
    "braid": "BRAID SHIELD",
}
_TITLE_SUFFIX = {"vi": " / ĐƠN ĐẶT HÀNG", "en": "", "ja": " / 注文書"}
_RFQ_TITLE_SUFFIX = {"vi": " / YÊU CẦU BÁO GIÁ", "en": "", "ja": " / 見積依頼書"}
_PO_SHEET = {"vi": "PO", "en": "PO", "ja": "注文書"}
_RFQ_SHEET = {"vi": "RFQ", "en": "RFQ", "ja": "見積依頼"}
_PRICE_FORMAT: dict[Currency, str] = {"USD": "0.0000", "JPY": "0.00", "VND": "#,##0"}
_AMOUNT_FORMAT: dict[Currency, str] = {"USD": "#,##0.00", "JPY": "#,##0", "VND": "#,##0"}
_AMOUNT_STEP: dict[Currency, Decimal] = {
    "USD": Decimal("0.01"),
    "JPY": Decimal(1),
    "VND": Decimal(1),
}
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)

# What these documents put in a cell: a subset of what openpyxl accepts.
_CellValue = str | int | Decimal | date | None


class StatedAttributes(BaseModel):
    """What a customer's description states about an item. Unset = not stated."""

    model_config = _FROZEN

    family: ItemFamily | None = None
    cores: int | None = Field(default=None, ge=1)
    gauge: str | None = None
    stranding: Stranding | None = None
    conductor: Conductor | None = None
    shield: Shield | None = None
    colour: Colour | None = None
    packaging: Packaging | None = None
    pack_length_m: Decimal | None = Field(default=None, gt=0)
    spec_no: str | None = None

    @classmethod
    def of_item(cls, item: Item) -> StatedAttributes:
        """Everything the catalogue knows about the item."""
        attributes = item.attributes
        return cls(
            family=item.family,
            cores=attributes.cores,
            gauge=attributes.gauge,
            stranding=attributes.stranding,
            conductor=attributes.conductor,
            shield=attributes.shield,
            colour=attributes.colour,
            packaging=attributes.packaging,
            pack_length_m=item.pack_multiple,
            spec_no=attributes.spec_no,
        )

    def describe(self) -> str:
        """The description the mock customers print, in a fixed word order.

        ``MULTI-CORE CABLE 2C 0.5mm2 STRANDED TINNED CU FOIL SHIELD BLACK REEL
        300M SPEC SP-5201``; whatever is not stated is left out.
        """
        words: list[str] = []
        if self.family is not None:
            words.append(_FAMILY_WORDS[self.family])
        if self.cores is not None:
            words.append(f"{self.cores}C")
        if self.gauge is not None:
            words.append(self.gauge)
        if self.stranding is not None:
            words.append(self.stranding.upper())
        if self.conductor is not None:
            words.append(_CONDUCTOR_WORDS[self.conductor])
        if self.shield is not None:
            words.append(_SHIELD_WORDS[self.shield])
        if self.colour is not None:
            words.append(self.colour.upper())
        if self.packaging is not None:
            words.append(self.packaging.upper())
        if self.pack_length_m is not None:
            words.append(f"{self.pack_length_m}M")
        if self.spec_no is not None:
            words.append(f"SPEC {self.spec_no}")
        return " ".join(words)


class FixtureSchedule(BaseModel):
    """One delivery of a line's quantity, as the PO prints it (ticket 21)."""

    model_config = _FROZEN

    requested_date: date
    quantity: Decimal = Field(gt=0)


class OrderLine(BaseModel):
    model_config = _FROZEN

    no: int = Field(ge=1)
    customer_item_code: str
    # None: described as the catalogue describes the item the code maps to.
    attributes: StatedAttributes | None = None
    # Printed exactly as given, for a description no catalogue would print
    # (ticket 06's formula test). Excludes ``attributes``.
    description: str | None = Field(default=None, min_length=1)
    quantity: Decimal = Field(gt=0)
    # The unit as the customer prints it; DW1's vocabulary has only ``m``.
    uom: str = Field(default="m", min_length=1, max_length=8)
    unit_price: Decimal = Field(gt=0)
    # The line's date; for a line split over several dates, the first one.
    requested_date: date
    # The dates and quantities a split line prints apart from the line, in
    # date order, as the document's `schedule_layout` says. They need not add
    # up to ``quantity``: a fixture may print a split that does not.
    schedules: tuple[FixtureSchedule, ...] = ()

    @model_validator(mode="after")
    def _described_one_way(self) -> Self:
        if self.attributes is not None and self.description is not None:
            raise ValueError(f"line {self.no}: attributes or a verbatim description, not both")
        return self

    @model_validator(mode="after")
    def _schedules_start_at_the_date(self) -> Self:
        days = [s.requested_date for s in self.schedules]
        if days and (len(days) < 2 or days != sorted(days) or days[0] != self.requested_date):
            raise ValueError(
                f"line {self.no}: schedules are two or more, in date order, from requested_date"
            )
        return self


Layout = Literal["xlsx", "xlsx_sheet_per_page", "pdf", "pdf_image"]

# What a customer's own form prints beside each label (ticket 30): the values
# a PO header shows, and the columns of its table (``note`` and ``ordered``,
# the PO date repeated on each line, map to no field).
FormHeaderField = Literal["po_no", "revision", "po_date", "currency", "supplier", "remarks"]
FormColumnField = Literal[
    "line_no",
    "customer_item_code",
    "description",
    "quantity",
    "uom",
    "unit_price",
    "amount",
    "requested_date",
    "note",
    "ordered",
]
_FORM_HEADER_NEEDS = frozenset({"po_no", "po_date", "currency"})
_FORM_LINE_FIELDS = frozenset(
    {
        "line_no",
        "customer_item_code",
        "description",
        "quantity",
        "uom",
        "unit_price",
        "amount",
        "requested_date",
    }
)


class PoForm(BaseModel):
    """A customer's own PO form, printed by their system (ticket 30): its
    title, the labels of its header in print order, the titles of its table's
    columns in print order, and the label of its total. A PDF form also says
    how its numbers group and how its dates read; a workbook keeps numbers and
    dates as numbers and dates."""

    model_config = _FROZEN

    title: str = Field(min_length=1)
    sheet: str = Field(default="PO", min_length=1, max_length=31)
    header: tuple[tuple[str, FormHeaderField], ...]
    columns: tuple[tuple[str, FormColumnField], ...]
    # None: the form prints no total row (an eval of a form no layout can take).
    total: str | None = Field(min_length=1)
    thousands: Literal[",", ".", " ", ""] = ","
    decimal: Literal[".", ","] = "."
    date_format: str = "%d/%m/%Y"

    @model_validator(mode="after")
    def _prints_what_a_po_needs(self) -> Self:
        fields = [field for _, field in self.header]
        if not set(fields) >= _FORM_HEADER_NEEDS or len(fields) != len(set(fields)):
            raise ValueError("a form prints the PO number, date and currency, once each")
        columns = [field for _, field in self.columns if field not in ("note", "ordered")]
        if set(columns) != _FORM_LINE_FIELDS or len(columns) != len(_FORM_LINE_FIELDS):
            raise ValueError("a form prints every line field in one column")
        if self.columns[0][1] != "line_no":
            raise ValueError("a form's first column numbers the lines")
        return self


# How a PO prints a line's schedules (ticket 21): rows under the line, each
# with a quantity and a date and nothing else, the line's own date cell left
# blank; or a "date: quantity; ..." list in the line's date cell.
ScheduleLayout = Literal["rows", "cell"]


class PurchaseOrderDocument(BaseModel):
    """One PO file, attached to message ``message_id`` under ``name``."""

    model_config = _FROZEN

    message_id: str
    name: str
    layout: Layout
    # The customer whose document this is; None for a buyer no customer is.
    customer_code: str | None = None
    # The buyer the document names when it is no customer (printed in English).
    buyer_name: str | None = None
    po_no: str
    revision: int = Field(ge=0)
    po_date: date
    currency: Currency
    remarks: str = ""
    # The intra-group layout: one worksheet per printed page.
    lines_per_page: int | None = Field(default=None, ge=1)
    # The PDF layout only: printed under the buyer's name, and as a second page.
    buyer_address: str = ""
    terms: tuple[str, ...] = ()
    # The TOTAL the buyer's system printed, when it is not the sum of the lines.
    printed_total: Decimal | None = None
    # The Excel layouts only: lines whose row is hidden, and pages (sheets) hidden.
    hidden_lines: tuple[int, ...] = ()
    hidden_pages: tuple[int, ...] = ()
    # The Excel layouts only: lines whose description cell is set in
    # `LEGACY_FONT`, a TCVN3 font (`legacy_encoding`).
    legacy_font_lines: tuple[int, ...] = ()
    # How the lines with schedules print them; set exactly when one has some.
    schedule_layout: ScheduleLayout | None = None
    # The customer's own form (ticket 30), for a PO no platform layout reads:
    # printed with its labels instead of the fixed ones, as an .xlsx or a PDF.
    form: PoForm | None = None
    lines: tuple[OrderLine, ...] = Field(min_length=1)

    @model_validator(mode="after")
    def _layout_has_what_it_needs(self) -> Self:
        if (self.customer_code is None) == (self.buyer_name is None):
            raise ValueError(f"{self.name}: a customer_code or a buyer_name, one of them")
        if (self.layout == "xlsx_sheet_per_page") != (self.lines_per_page is not None):
            raise ValueError(
                f"{self.name}: lines_per_page goes with the sheet-per-page layout only"
            )
        # The Excel layouts have no place for them, and a field the file does
        # not print is a fixture edit that silently changes nothing.
        if self.layout != "pdf" and (self.buyer_address or self.terms):
            raise ValueError(f"{self.name}: buyer_address and terms are printed by the PDF only")
        if self.layout not in ("xlsx", "xlsx_sheet_per_page") and (
            self.hidden_lines or self.hidden_pages or self.legacy_font_lines
        ):
            raise ValueError(f"{self.name}: only an Excel layout hides rows or sets cell fonts")
        numbers = [line.no for line in self.lines]
        # Increasing, so a README line number names one printed line. Gaps are
        # allowed: a PO that skips a number is a fixture (`line_total_mismatch`).
        if numbers != sorted(set(numbers)):
            raise ValueError(f"{self.name}: lines are numbered in increasing order")
        if not set(self.hidden_lines) <= set(numbers):
            raise ValueError(f"{self.name}: hidden_lines names a line the PO lacks")
        if not set(self.legacy_font_lines) <= set(numbers):
            raise ValueError(f"{self.name}: legacy_font_lines names a line the PO lacks")
        if (self.schedule_layout is None) == any(line.schedules for line in self.lines):
            raise ValueError(f"{self.name}: schedule_layout goes with lines that have schedules")
        if self.schedule_layout == "rows" and self.layout not in ("xlsx", "xlsx_sheet_per_page"):
            raise ValueError(f"{self.name}: only an Excel layout prints schedule rows")
        if self.form is not None and (
            self.layout not in ("xlsx", "pdf") or self.schedule_layout or self.lines_per_page
        ):
            raise ValueError(f"{self.name}: a customer's form is one .xlsx sheet or one PDF page")
        pages = len(self.pages())
        if not all(1 <= page < pages for page in self.hidden_pages):
            # The last page carries the TOTAL, which stays visible.
            raise ValueError(f"{self.name}: hidden_pages names no page before the last")
        return self

    def pages(self) -> list[tuple[OrderLine, ...]]:
        """The lines as the file prints them, page by page (sheet by sheet)."""
        size = self.lines_per_page or len(self.lines)
        return [self.lines[start : start + size] for start in range(0, len(self.lines), size)]


class QuoteRequestLine(BaseModel):
    model_config = _FROZEN

    no: int = Field(ge=1)
    customer_item_code: str
    attributes: StatedAttributes | None = None
    # None: the request leaves the quantity blank (`rfq_incomplete`).
    quantity: Decimal | None = Field(default=None, gt=0)
    target_price: Decimal | None = Field(default=None, gt=0)
    required_date: date


class QuoteRequestDocument(BaseModel):
    """One RFQ file, attached to message ``message_id`` under ``name``."""

    model_config = _FROZEN

    message_id: str
    name: str
    customer_code: str
    rfq_no: str
    rfq_date: date
    quote_due: date
    currency: Currency
    remarks: str = ""
    lines: tuple[QuoteRequestLine, ...] = Field(min_length=1)


class FreeRequestLine(BaseModel):
    model_config = _FROZEN

    no: int = Field(ge=1)
    customer_item_code: str | None = None
    description: str
    # None: the request leaves it blank (`rfq_incomplete`).
    quantity: Decimal | None = Field(default=None, gt=0)
    required_date: date | None = None
    target_price: Decimal | None = Field(default=None, gt=0)
    # A spreadsheet row the customer hid: not part of the request a person
    # sees, and not sent to a model (`readers.text_parts`).
    hidden: bool = False


class FreeRequestDocument(BaseModel):
    """A request for quotation in the customer's own form (ticket 29): the
    mail body of message ``message_id``, or an attachment ``name`` printed as
    a text PDF or the customer's own spreadsheet. No reader reads it; where a
    tenant turns it on, a model does, and this is what it states."""

    model_config = _FROZEN

    message_id: str
    form: Literal["body", "pdf", "xlsx"]
    name: str | None
    customer_code: str
    rfq_no: str | None
    quote_due: date | None
    currency: Currency
    lines: tuple[FreeRequestLine, ...] = Field(min_length=1)

    @model_validator(mode="after")
    def _a_file_has_a_name(self) -> Self:
        if (self.form == "body") != (self.name is None):
            raise ValueError(f"{self.message_id}: an attachment is named, a body is not")
        return self


class DesignReplyLine(BaseModel):
    model_config = _FROZEN

    # The request's line this answers.
    no: int = Field(ge=1)
    bp_code: str
    spec_no: str
    # None while Design has still to create the item's code.
    prv_code: str | None = None
    copper_kg_per_km: Decimal = Field(gt=0)


class DesignReplyDocument(BaseModel):
    """Design's answer to one YCBG, attached to message ``message_id``.

    Names the YCBG and nothing else of the request: matching the reply to a
    case is by that number only.
    """

    model_config = _FROZEN

    message_id: str
    name: str
    ycbg_no: str
    reply_date: date
    lines: tuple[DesignReplyLine, ...] = Field(min_length=1)


class _Buyer(NamedTuple):
    name: str
    language: Language


class _Catalogue:
    """The slice of master data the documents print: names and descriptions."""

    def __init__(self, data_dir: Path) -> None:
        self.customers = {c.code: c for c in read_records(data_dir / "customers.json", Customer)}
        self.items = {i.prv_code: i for i in read_records(data_dir / "items.json", Item)}
        self.convert = {
            (e.customer_code, e.customer_item_code): e.prv_code
            for e in read_records(data_dir / "convert_list.json", ConvertEntry)
        }

    def buyer(self, po: PurchaseOrderDocument) -> _Buyer:
        if po.customer_code is None:
            assert po.buyer_name is not None  # the document's validator
            return _Buyer(po.buyer_name, "en")
        customer = self.customers[po.customer_code]
        return _Buyer(customer.name, customer.language)

    def describe(
        self,
        customer_code: str | None,
        customer_item_code: str,
        stated: StatedAttributes | None,
        verbatim: str | None = None,
    ) -> str:
        if verbatim is not None:
            return verbatim
        if stated is not None:
            return stated.describe()
        prv_code = self.convert.get((customer_code, customer_item_code)) if customer_code else None
        if prv_code is None:
            raise ValueError(
                f"{customer_code} line {customer_item_code} states no attributes and has no"
                " convert entry, so there is nothing to describe it by"
            )
        return StatedAttributes.of_item(self.items[prv_code]).describe()


def load_purchase_orders(data_dir: Path = DATA_DIR) -> tuple[PurchaseOrderDocument, ...]:
    return read_records(data_dir / "purchase_orders.json", PurchaseOrderDocument)


def load_quote_requests(data_dir: Path = DATA_DIR) -> tuple[QuoteRequestDocument, ...]:
    return read_records(data_dir / "quote_requests.json", QuoteRequestDocument)


def load_design_replies(data_dir: Path = DATA_DIR) -> tuple[DesignReplyDocument, ...]:
    return read_records(data_dir / "design_replies.json", DesignReplyDocument)


def load_free_requests(data_dir: Path = DATA_DIR) -> tuple[FreeRequestDocument, ...]:
    return read_records(data_dir / "free_requests.json", FreeRequestDocument)


def render_attachments(data_dir: Path = DATA_DIR) -> dict[str, bytes]:
    """Every attachment's bytes, by its file name in `attachments/`."""
    catalogue = _Catalogue(data_dir)
    rendered: dict[str, bytes] = {}
    for po in load_purchase_orders(data_dir):
        rendered[attachment_file_name(po.message_id, po.name)] = render_po(po, catalogue)
    for rfq in load_quote_requests(data_dir):
        customer = catalogue.customers[rfq.customer_code]
        rendered[attachment_file_name(rfq.message_id, rfq.name)] = _rfq_xlsx(
            rfq, customer, catalogue
        )
    for reply in load_design_replies(data_dir):
        rendered[attachment_file_name(reply.message_id, reply.name)] = _design_reply_xlsx(reply)
    for free in load_free_requests(data_dir):
        customer = catalogue.customers[free.customer_code]
        if free.form == "pdf":
            data = _free_request_pdf(free, customer)
        elif free.form == "xlsx":
            data = _free_request_xlsx(free, customer)
        else:
            continue
        assert free.name is not None  # the document's validator
        rendered[attachment_file_name(free.message_id, free.name)] = data
    return rendered


def render_po(po: PurchaseOrderDocument, catalogue: _Catalogue) -> bytes:
    """One PO file's bytes: the fixed layouts, a scan, or the customer's form."""
    if po.form is not None:
        if po.layout == "pdf":
            return _form_pdf(po, po.form, catalogue)
        return _form_xlsx(po, po.form, catalogue)
    match po.layout:
        case "pdf":
            return _po_pdf(po, catalogue)
        case "pdf_image":
            return _scanned_pdf(po)
        case "xlsx" | "xlsx_sheet_per_page":
            return _po_xlsx(po, catalogue)


def catalogue_of(data_dir: Path = DATA_DIR) -> _Catalogue:
    """The master data a document prints its names and descriptions from."""
    return _Catalogue(data_dir)


def main() -> None:
    rendered = render_attachments()
    ATTACHMENTS_DIR.mkdir(exist_ok=True)
    # The directory holds exactly what the fixtures produce. A file left from
    # a renamed or removed attachment would fail the sync test however often
    # this ran.
    for stale in sorted(ATTACHMENTS_DIR.iterdir()):
        if stale.is_file() and stale.name not in rendered:
            stale.unlink()
            print(f"removed {stale.name}")
    for file_name, data in rendered.items():
        (ATTACHMENTS_DIR / file_name).write_bytes(data)
        print(f"wrote {file_name} ({len(data)} bytes)")


# ------------------------------------------------------------------- xlsx --


def _amount(quantity: Decimal, unit_price: Decimal, currency: Currency) -> Decimal:
    return (quantity * unit_price).quantize(_AMOUNT_STEP[currency], ROUND_HALF_UP)


def _total(po: PurchaseOrderDocument) -> Decimal:
    """The TOTAL as printed: the buyer's sum of the line amounts, unless the
    fixture says the buyer's system printed another."""
    if po.printed_total is not None:
        return po.printed_total
    return sum(
        (_amount(line.quantity, line.unit_price, po.currency) for line in po.lines), Decimal(0)
    )


def _set(sheet: Worksheet, row: int, column: int, value: _CellValue) -> None:
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        # Text the customer typed, not a formula: openpyxl would otherwise
        # store it as one, with no cached value, which is a different fixture.
        cell.data_type = "s"


def _header(sheet: Worksheet, rows: Sequence[tuple[str, _CellValue, str, _CellValue]]) -> None:
    """Label/value pairs from row 4, two pairs per row (columns A-B and D-E)."""
    for offset, (label, value, second_label, second_value) in enumerate(rows):
        row = 4 + offset
        sheet.cell(row=row, column=1, value=label)
        cell = sheet.cell(row=row, column=2, value=value)
        if isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
        if second_label:
            sheet.cell(row=row, column=4, value=second_label)
            second = sheet.cell(row=row, column=5, value=second_value)
            if isinstance(second_value, date):
                second.number_format = "yyyy-mm-dd"


def _table(
    sheet: Worksheet,
    columns: Sequence[str],
    rows: Sequence[Sequence[_CellValue]],
    formats: Sequence[str | None],
) -> int:
    """A header on row 9 and one row per entry; returns the next free row."""
    for column, title in enumerate(columns, start=1):
        sheet.cell(row=9, column=column, value=title)
    for offset, values in enumerate(rows):
        for column, (value, number_format) in enumerate(zip(values, formats, strict=True), start=1):
            _set(sheet, 10 + offset, column, value)
            if number_format is not None:
                sheet.cell(row=10 + offset, column=column).number_format = number_format
    for letter, width in zip("ABCDEFGH", (6, 18, 72, 12, 6, 13, 15, 15), strict=False):
        sheet.column_dimensions[letter].width = width
    return 10 + len(rows)


def _po_xlsx(po: PurchaseOrderDocument, catalogue: _Catalogue) -> bytes:
    buyer = catalogue.buyer(po)
    workbook = _new_workbook()
    pages = po.pages()
    for number, lines in enumerate(pages, start=1):
        title = _PO_SHEET[buyer.language] if po.lines_per_page is None else f"Page {number}"
        sheet: Worksheet = workbook.create_sheet(title)
        if number in po.hidden_pages:
            sheet.sheet_state = "hidden"
        sheet["A1"] = buyer.name.upper()
        sheet["A2"] = "PURCHASE ORDER" + _TITLE_SUFFIX[buyer.language]
        page_label = ("Page", f"{number} / {len(pages)}") if po.lines_per_page else ("", "")
        _header(
            sheet,
            [
                ("PO No.", po.po_no, "Revision", po.revision),
                ("PO Date", po.po_date, "Currency", po.currency),
                ("Supplier", SELLER_NAME, *page_label),
                ("Remarks", po.remarks, "", ""),
            ],
        )
        rows: list[list[_CellValue]] = []
        # The row each line starts on, from row 10.
        first_rows: dict[int, int] = {}
        for line in lines:
            first_rows[line.no] = 10 + len(rows)
            rows.append(
                [
                    line.no,
                    line.customer_item_code,
                    catalogue.describe(
                        po.customer_code, line.customer_item_code, line.attributes, line.description
                    ),
                    line.quantity,
                    line.uom,
                    line.unit_price,
                    _amount(line.quantity, line.unit_price, po.currency),
                    _excel_date_cell(po, line),
                ]
            )
            if line.schedules and po.schedule_layout == "rows":
                rows += [
                    [None, None, None, s.quantity, None, None, None, s.requested_date]
                    for s in line.schedules
                ]
        next_row = _table(
            sheet,
            (
                "No.",
                "Customer Part No.",
                "Description",
                "Quantity",
                "UoM",
                "Unit Price",
                "Amount",
                "Requested Date",
            ),
            rows,
            (
                None,
                None,
                None,
                "#,##0",
                None,
                _PRICE_FORMAT[po.currency],
                _AMOUNT_FORMAT[po.currency],
                "yyyy-mm-dd",
            ),
        )
        for line in lines:
            if line.no in po.hidden_lines:
                sheet.row_dimensions[first_rows[line.no]].hidden = True
            if line.no in po.legacy_font_lines:
                sheet.cell(row=first_rows[line.no], column=3).font = Font(name=LEGACY_FONT)
        if number == len(pages):
            sheet.cell(row=next_row, column=6, value="TOTAL")
            total_cell = sheet.cell(row=next_row, column=7, value=_total(po))
            total_cell.number_format = _AMOUNT_FORMAT[po.currency]
    return _xlsx_bytes(workbook, creator=buyer.name, stamp=po.po_date)


def _excel_date_cell(po: PurchaseOrderDocument, line: OrderLine) -> _CellValue:
    """A line's date cell: its date; blank above schedule rows; or the
    schedules as ``15/11/2026: 6100; 01/12/2026: 6100``."""
    if not line.schedules:
        return line.requested_date
    if po.schedule_layout == "rows":
        return None
    return "; ".join(f"{_vi_date(s.requested_date)}: {s.quantity}" for s in line.schedules)


def _rfq_xlsx(rfq: QuoteRequestDocument, customer: Customer, catalogue: _Catalogue) -> bytes:
    workbook = _new_workbook()
    sheet: Worksheet = workbook.create_sheet(_RFQ_SHEET[customer.language])
    sheet["A1"] = customer.name.upper()
    sheet["A2"] = "REQUEST FOR QUOTATION" + _RFQ_TITLE_SUFFIX[customer.language]
    _header(
        sheet,
        [
            ("RFQ No.", rfq.rfq_no, "RFQ Date", rfq.rfq_date),
            ("Quote Due", rfq.quote_due, "Currency", rfq.currency),
            ("To", SELLER_NAME, "", ""),
            ("Remarks", rfq.remarks, "", ""),
        ],
    )
    rows: list[list[_CellValue]] = [
        [
            line.no,
            line.customer_item_code,
            catalogue.describe(rfq.customer_code, line.customer_item_code, line.attributes),
            line.quantity,
            "m",
            line.target_price,
            line.required_date,
        ]
        for line in rfq.lines
    ]
    _table(
        sheet,
        (
            "No.",
            "Customer Part No.",
            "Description",
            "Quantity",
            "UoM",
            "Target Price",
            "Required Date",
        ),
        rows,
        (None, None, None, "#,##0", None, _PRICE_FORMAT[rfq.currency], "yyyy-mm-dd"),
    )
    return _xlsx_bytes(workbook, creator=customer.name, stamp=rfq.rfq_date)


def _free_request_xlsx(request: FreeRequestDocument, customer: Customer) -> bytes:
    """The customer's own Japanese request sheet: no "RFQ No." label, its own
    column titles, and a hidden row (an instruction a person would not see)."""
    workbook = _new_workbook()
    sheet: Worksheet = workbook.create_sheet("見積依頼")
    sheet["A1"] = f"{customer.name} 見積依頼書"
    sheet["A3"], sheet["B3"] = "依頼番号", request.rfq_no
    sheet["A4"], sheet["B4"] = "回答期限", request.quote_due
    sheet["B4"].number_format = "yyyy/mm/dd"
    sheet["A5"], sheet["B5"] = "通貨", request.currency
    for column, title in enumerate(("No", "品番", "品名", "数量", "単位", "希望納期"), start=1):
        sheet.cell(row=7, column=column, value=title).font = Font(bold=True)
    for offset, line in enumerate(request.lines):
        row = 8 + offset
        values: tuple[_CellValue, ...] = (
            line.no,
            line.customer_item_code,
            line.description,
            line.quantity,
            "m",
            line.required_date,
        )
        for column, value in enumerate(values, start=1):
            _set(sheet, row, column, value)
        sheet.cell(row=row, column=4).number_format = "#,##0"
        sheet.cell(row=row, column=6).number_format = "yyyy/mm/dd"
        if line.hidden:
            sheet.row_dimensions[row].hidden = True
    return _xlsx_bytes(workbook, creator=customer.name, stamp=date(2026, 10, 2))


def _free_request_pdf(request: FreeRequestDocument, customer: Customer) -> bytes:
    """The customer's own English request, printed as a text PDF: a header,
    the items as text rows, and a line on reels that is no quantity."""
    pdfmetrics.registerFont(TTFont(_FONT, str(FONT_PATH)))
    buffer = io.BytesIO()
    width, height = A4
    canvas = Canvas(buffer, pagesize=(width, height), invariant=1, pageCompression=0)
    canvas.setTitle(f"Request for quotation {request.rfq_no}")
    canvas.setAuthor(customer.name)

    def text(x: float, y: float, value: str, size: float = 9) -> None:
        canvas.setFont(_FONT, size)
        canvas.drawString(x, y, value)

    y = height - 50
    text(40, y, customer.name.upper(), 11)
    y -= 26
    text(40, y, "REQUEST FOR QUOTATION", 14)
    due = request.quote_due.strftime("%d %b %Y") if request.quote_due else "-"
    for line in (
        f"RFQ No.: {request.rfq_no}",
        f"To: {SELLER_NAME}",
        f"Please quote in {request.currency} and reply by {due}.",
    ):
        y -= 16
        text(40, y, line)
    y -= 26
    text(40, y, "Item / Our part no. / Description / Quantity / Required by / Target price", 8)
    for item in request.lines:
        y -= 16
        quantity = f"{item.quantity:,.0f} m" if item.quantity is not None else "-"
        needed = item.required_date.strftime("%d %b %Y") if item.required_date else "-"
        target = f"USD {item.target_price}" if item.target_price is not None else "-"
        text(
            40,
            y,
            f"{item.no}. {item.customer_item_code} {item.description} - {quantity}"
            f" - by {needed} - target {target}",
            8,
        )
    y -= 24
    text(40, y, "Packing: 305 m reels preferred where the item allows.", 8)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def _design_reply_xlsx(reply: DesignReplyDocument) -> bytes:
    """The seller's own Design reply form: the YCBG it answers and, per
    requested line, what Design specified."""
    workbook = _new_workbook()
    sheet: Worksheet = workbook.create_sheet("YCBG")
    sheet["A1"] = SELLER_NAME.upper() + " / DESIGN"
    sheet["A2"] = "YCBG REPLY / PHẢN HỒI YCBG"
    _header(
        sheet,
        [
            ("YCBG No.", reply.ycbg_no, "Reply Date", reply.reply_date),
            ("From", "Design", "", ""),
        ],
    )
    rows: list[list[_CellValue]] = [
        [line.no, line.bp_code, line.spec_no, line.prv_code, line.copper_kg_per_km]
        for line in reply.lines
    ]
    _table(
        sheet,
        ("No.", "BP Code", "Spec No.", "PRV Code", "Copper (kg/km)"),
        rows,
        (None, None, None, None, "0.00"),
    )
    return _xlsx_bytes(workbook, creator=SELLER_NAME, stamp=reply.reply_date)


def _new_workbook() -> Workbook:
    """A workbook without the default sheet: each document names its own."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    return workbook


def _xlsx_bytes(workbook: Workbook, *, creator: str, stamp: date) -> bytes:
    """The workbook as a file whose bytes depend on its content only.

    `ExcelWriter` directly rather than `Workbook.save`, which stamps the
    current time into the document properties.
    """
    if not LXML:
        # Measured: without lxml every workbook serialises differently.
        raise RuntimeError(
            "openpyxl is not writing through lxml, so the workbooks would differ from the"
            " committed ones: run in the dev environment (uv sync --all-packages)"
        )
    moment = datetime(stamp.year, stamp.month, stamp.day)
    workbook.properties.creator = creator
    workbook.properties.created = moment
    workbook.properties.modified = moment
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as archive:
        ExcelWriter(workbook, archive).save()
    return _fixed_zip(buffer.getvalue())


def _fixed_zip(data: bytes) -> bytes:
    """The same archive with every entry's time and origin fixed.

    `ZipFile.writestr` stamps entries with the current time and the host
    system (Windows or Unix), and deflate output varies with the zlib build;
    any of those would make two machines produce different bytes. Stored, not
    deflated: the files are small and Excel reads either.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_EPOCH)
            entry.create_system = 3
            entry.external_attr = 0o644 << 16
            target.writestr(entry, source.read(info.filename))
    return buffer.getvalue()


# --------------------------------------------------- a customer's form --


def _form_values(
    po: PurchaseOrderDocument, line: OrderLine, catalogue: _Catalogue
) -> dict[str, _CellValue]:
    return {
        "line_no": line.no,
        "customer_item_code": line.customer_item_code,
        "description": catalogue.describe(
            po.customer_code, line.customer_item_code, line.attributes, line.description
        ),
        "quantity": line.quantity,
        "uom": line.uom,
        "unit_price": line.unit_price,
        "amount": _amount(line.quantity, line.unit_price, po.currency),
        "requested_date": line.requested_date,
        "note": None,
        "ordered": po.po_date,
    }


def _form_header(po: PurchaseOrderDocument, field: FormHeaderField) -> _CellValue:
    values: dict[str, _CellValue] = {
        "po_no": po.po_no,
        "revision": po.revision,
        "po_date": po.po_date,
        "currency": po.currency,
        "supplier": SELLER_NAME,
        "remarks": po.remarks,
    }
    return values[field]


def _form_xlsx(po: PurchaseOrderDocument, form: PoForm, catalogue: _Catalogue) -> bytes:
    """The form as one sheet: the buyer and the title, one header label per
    row with its value beside it, the table under it, and the total row."""
    buyer = catalogue.buyer(po)
    workbook = _new_workbook()
    sheet: Worksheet = workbook.create_sheet(form.sheet)
    sheet["A1"] = buyer.name.upper()
    sheet["A2"] = form.title
    for offset, (label, field) in enumerate(form.header):
        sheet.cell(row=4 + offset, column=1, value=label)
        value = _form_header(po, field)
        cell = sheet.cell(row=4 + offset, column=2, value=value)
        if isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
    first = 4 + len(form.header) + 1
    formats = {
        "quantity": "#,##0",
        "unit_price": _PRICE_FORMAT[po.currency],
        "amount": _AMOUNT_FORMAT[po.currency],
        "requested_date": "yyyy-mm-dd",
        "ordered": "yyyy-mm-dd",
    }
    for column, (title, _) in enumerate(form.columns, start=1):
        sheet.cell(row=first, column=column, value=title)
    for offset, line in enumerate(po.lines, start=1):
        values = _form_values(po, line, catalogue)
        for column, (_, printed) in enumerate(form.columns, start=1):
            _set(sheet, first + offset, column, values[printed])
            if printed in formats:
                sheet.cell(row=first + offset, column=column).number_format = formats[printed]
    if form.total is not None:
        total_row = first + len(po.lines) + 1
        amount_column = next(i for i, (_, f) in enumerate(form.columns, start=1) if f == "amount")
        sheet.cell(row=total_row, column=max(amount_column - 1, 1), value=form.total)
        total = sheet.cell(row=total_row, column=amount_column, value=_total(po))
        total.number_format = _AMOUNT_FORMAT[po.currency]
    return _xlsx_bytes(workbook, creator=buyer.name, stamp=po.po_date)


def _grouped(value: Decimal, thousands: str, decimal: str) -> str:
    whole, _, fraction = f"{value:f}".partition(".")
    grouped = f"{int(whole):,}".replace(",", thousands)
    return f"{grouped}{decimal}{fraction}" if fraction else grouped


def _form_pdf(po: PurchaseOrderDocument, form: PoForm, catalogue: _Catalogue) -> bytes:
    """The form as one landscape page: ``Label: value`` header lines, the
    table's titles and rows on columns wide apart, and ``Total: amount CUR``."""
    buyer = catalogue.buyer(po)
    pdfmetrics.registerFont(TTFont(_FONT, str(FONT_PATH)))
    buffer = io.BytesIO()
    width, height = landscape(A4)
    canvas = Canvas(buffer, pagesize=(width, height), invariant=1, pageCompression=0)
    canvas.setTitle(f"{form.title} {po.po_no}")
    canvas.setAuthor(buyer.name)

    def text(x: float, y: float, value: str, size: float = 9) -> None:
        canvas.setFont(_FONT, size)
        canvas.drawString(x, y, value)

    def shown(value: _CellValue) -> str:
        if isinstance(value, Decimal):
            return _grouped(value, form.thousands, form.decimal)
        if isinstance(value, date):
            return value.strftime(form.date_format)
        return "" if value is None else str(value)

    y = height - 50
    text(40, y, buyer.name.upper(), 12)
    y -= 26
    text(40, y, form.title, 14)
    for label, field in form.header:
        y -= 15
        text(40, y, f"{label}: {shown(_form_header(po, field))}")
    # Wide enough apart that the text layer keeps two spaces between cells.
    widths = {
        "line_no": 24,
        "customer_item_code": 80,
        "description": 400,
        "quantity": 46,
        "uom": 24,
        "unit_price": 50,
        "amount": 60,
        "requested_date": 58,
        "note": 36,
        "ordered": 58,
    }
    xs: list[float] = []
    x = 40.0
    for title, printed in form.columns:
        xs.append(x)
        # A title wider than its column's values pushes the next column on.
        x += max(widths[printed], pdfmetrics.stringWidth(title, _FONT, 8) + 14)
    y -= 26
    for (title, _), at in zip(form.columns, xs, strict=True):
        text(at, y, title, 8)
    for line in po.lines:
        y -= 14
        values = _form_values(po, line, catalogue)
        for (_, printed), at in zip(form.columns, xs, strict=True):
            cell = shown(values[printed]) or "-"
            text(at, y, cell, 7.5)
    if form.total is not None:
        y -= 22
        text(40, y, f"{form.total}: {shown(_total(po))} {po.currency}", 9)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()


# -------------------------------------------------------------------- pdf --


def _vi_number(value: Decimal) -> str:
    """Vietnamese grouping: ``.`` between thousands, ``,`` before decimals."""
    return f"{value:,}".replace(",", " ").replace(".", ",").replace(" ", ".")


def _vi_date(day: date) -> str:
    return day.strftime("%d/%m/%Y")


# x of each column; numbers are right-aligned on theirs.
_PDF_COLUMNS: tuple[tuple[str, float, bool], ...] = (
    ("STT", 55, True),
    ("Mã hàng KH", 62, False),
    ("Mô tả hàng hóa", 140, False),
    ("Số lượng", 612, True),
    ("ĐVT", 618, False),
    ("Đơn giá", 690, True),
    ("Thành tiền", 770, True),
    ("Ngày giao", 778, False),
)
# The same columns with room for a "date: quantity" list in the last one.
_PDF_LIST_COLUMNS: tuple[tuple[str, float, bool], ...] = (
    ("STT", 55, True),
    ("Mã hàng KH", 62, False),
    ("Mô tả hàng hóa", 140, False),
    ("Số lượng", 400, True),
    ("ĐVT", 406, False),
    ("Đơn giá", 470, True),
    ("Thành tiền", 550, True),
    ("Ngày giao", 560, False),
)


def _pdf_date_cell(line: OrderLine) -> str:
    """``16/11/2026: 3.050; 30/11/2026: 3.050`` for a line with schedules."""
    if not line.schedules:
        return _vi_date(line.requested_date)
    return "; ".join(
        f"{_vi_date(s.requested_date)}: {_vi_number(s.quantity)}" for s in line.schedules
    )


def _po_pdf(po: PurchaseOrderDocument, catalogue: _Catalogue) -> bytes:
    """A Vietnamese text PDF: header and lines on page 1, terms on page 2."""
    buyer = catalogue.buyer(po)
    if buyer.language != "vi":
        raise ValueError(f"{po.name}: the PDF layout is the Vietnamese one")
    pdfmetrics.registerFont(TTFont(_FONT, str(FONT_PATH)))
    buffer = io.BytesIO()
    width, height = landscape(A4)
    canvas = Canvas(buffer, pagesize=(width, height), invariant=1, pageCompression=0)
    canvas.setTitle(f"Đơn đặt hàng {po.po_no}")
    canvas.setAuthor(buyer.name)
    pages = 2 if po.terms else 1

    def text(x: float, y: float, value: str, size: float = 9, right: bool = False) -> None:
        canvas.setFont(_FONT, size)
        if right:
            canvas.drawRightString(x, y, value)
        else:
            canvas.drawString(x, y, value)

    y = height - 50
    text(40, y, buyer.name.upper(), 12)
    if po.buyer_address:
        y -= 14
        text(40, y, po.buyer_address, 8)
    y -= 28
    text(40, y, "ĐƠN ĐẶT HÀNG (PURCHASE ORDER)", 14)
    for label, value in (
        ("Số đơn hàng", po.po_no),
        ("Lần sửa đổi", str(po.revision)),
        ("Ngày đặt hàng", _vi_date(po.po_date)),
        ("Nhà cung cấp", SELLER_NAME),
        ("Đơn vị tiền tệ", po.currency),
    ):
        y -= 15
        text(40, y, f"{label}: {value}")

    # A "date: quantity" list needs a wide date column.
    columns = _PDF_LIST_COLUMNS if po.schedule_layout == "cell" else _PDF_COLUMNS
    y -= 26
    for title, x, right in columns:
        text(x, y, title, 8, right)
    for line in po.lines:
        y -= 14
        values = (
            str(line.no),
            line.customer_item_code,
            catalogue.describe(
                po.customer_code, line.customer_item_code, line.attributes, line.description
            ),
            _vi_number(line.quantity),
            line.uom,
            _vi_number(line.unit_price),
            _vi_number(_amount(line.quantity, line.unit_price, po.currency)),
            _pdf_date_cell(line),
        )
        for (_, x, right), value in zip(columns, values, strict=True):
            text(x, y, value, 7.5, right)
    y -= 22
    text(770, y, f"Tổng cộng: {_vi_number(_total(po))} {po.currency}", 9, right=True)
    if po.remarks:
        y -= 22
        text(40, y, f"Ghi chú: {po.remarks}", 8)
    text(width - 40, 30, f"Trang 1/{pages}", 8, right=True)
    canvas.showPage()
    if po.terms:
        y = height - 50
        text(40, y, "ĐIỀU KHOẢN CHUNG", 12)
        for number, term in enumerate(po.terms, start=1):
            y -= 18
            text(40, y, f"{number}. {term}", 9)
        text(width - 40, 30, f"Trang 2/{pages}", 8, right=True)
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


# A4 portrait in points, drawn at one pixel per point.
_SCAN_WIDTH, _SCAN_HEIGHT = 595, 842


def _scanned_pdf(po: PurchaseOrderDocument) -> bytes:
    """A one-page PDF holding a single 1-bit image and no text at all.

    What a PO scanned on an office copier looks like to a parser: there is
    nothing to extract, which is the point of the fixture. The page is drawn
    as bars (a letterhead, header rows, one table row per line) rather than as
    glyphs, so no font rasteriser can make two machines draw it differently,
    and the PDF is written by hand, uncompressed, because image compression
    would go through the platform's zlib.
    """
    pixels = [[True] * _SCAN_WIDTH for _ in range(_SCAN_HEIGHT)]  # True is white

    def bar(x: int, y: int, w: int, h: int) -> None:
        for pixel_row in pixels[y : y + h]:
            pixel_row[x : x + w] = [False] * min(w, _SCAN_WIDTH - x)

    bar(40, 40, 260, 14)  # the buyer's name
    bar(40, 70, 180, 10)  # PURCHASE ORDER
    for pair in range(5):  # header label/value pairs
        bar(40, 100 + 16 * pair, 70, 6)
        bar(130, 100 + 16 * pair, 40 + 23 * ((pair * 7 + len(po.po_no)) % 5), 6)
    table_top = 200
    for index, _line in enumerate(po.lines):
        y = table_top + 22 * index
        bar(40, y, 515, 1)  # the row's rule
        bar(46, y + 7, 10, 6)  # No.
        bar(66, y + 7, 70, 6)  # code
        bar(146, y + 7, 150 + 17 * ((index * 5 + 3) % 7), 6)  # description
        bar(420, y + 7, 45, 6)  # quantity
        bar(480, y + 7, 60, 6)  # amount
    bar(40, table_top + 22 * len(po.lines), 515, 1)
    bar(380, table_top + 22 * len(po.lines) + 30, 160, 10)  # TOTAL
    bar(420, 720, 110, 2)  # the signature line

    row_bytes = (_SCAN_WIDTH + 7) // 8
    image = bytearray()
    for pixel_row in pixels:
        packed = bytearray(row_bytes)
        for x, white in enumerate(pixel_row):
            if white:
                packed[x // 8] |= 0x80 >> (x % 8)
        image += packed
    content = f"q {_SCAN_WIDTH} 0 0 {_SCAN_HEIGHT} 0 0 cm /Scan Do Q\n".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {_SCAN_WIDTH} {_SCAN_HEIGHT}]"
            " /Resources << /XObject << /Scan 4 0 R >> >> /Contents 5 0 R >>"
        ).encode("ascii"),
        (
            f"<< /Type /XObject /Subtype /Image /Width {_SCAN_WIDTH} /Height {_SCAN_HEIGHT}"
            f" /ColorSpace /DeviceGray /BitsPerComponent 1 /Length {len(image)} >>\nstream\n"
        ).encode("ascii")
        + bytes(image)
        + b"\nendstream",
        f"<< /Length {len(content)} >>\nstream\n".encode("ascii") + content + b"endstream",
    ]
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii")
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


if __name__ == "__main__":
    main()
