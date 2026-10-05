"""`RfqReaderPort` and `DesignReplyReaderPort` for Excel files, every value with its cell.

The layouts are the ones the mock customers and the seller's Design form print
(`adapters/mock/README.md`, "Reading the attachments"): labels with their value
in the cell to the right, above a table whose header row names its columns.
Each value is taken from the cell its label or column names, and from nowhere
else. A remark or a description that reads like an instruction is text in a
cell; nothing here interprets it.

The file is attacker-controlled and is parsed inside the API process, so the
work is bounded before openpyxl sees it: by what the archive declares it holds,
and by the rows and columns read, whatever dimensions a sheet claims. A
macro-enabled workbook is not read at all (spec decision 13), as on the order
side.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable, Mapping
from datetime import date, datetime
from decimal import Decimal
from typing import Final

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from dw_sales.adapters.readers.excel import MACROS_PART
from dw_sales.application.quote_ports import QuoteFileUnreadableError
from dw_sales.domain.anchors import SourceAnchor
from dw_sales.domain.catalog import Currency, CustomerItemCode, DocumentNo, Name, PrvCode, Uom
from dw_sales.domain.messages import Attachment, AttachmentContent
from dw_sales.domain.quotes import (
    BpCode,
    CopperWeight,
    Description,
    DesignReplyDocument,
    DesignReplyLine,
    Quantity,
    RequestedItem,
    RfqDocument,
    Sourced,
    UnitPrice,
)

# Stamped on every request it reads (`RfqDocument.parser_version`).
PARSER_VERSION: Final = "excel_rfq_reader@1.0.0"
_SUFFIX: Final = ".xlsx"
# A request for quotation is a few kilobytes. The bound is on what the archive
# declares it holds uncompressed, which zipfile holds an entry to as it inflates.
_MAX_UNCOMPRESSED_BYTES: Final = 16 * 1024 * 1024
_MAX_ENTRIES: Final = 1000
# What is read of each sheet, whatever dimensions it declares.
_MAX_ROWS: Final = 1000
_MAX_COLUMNS: Final = 26

# Where the buyer's name is printed, in capitals, on the customers' forms.
_BUYER_AT: Final = (1, 1)

_RFQ_NO: Final = "RFQ No."
_RFQ_DATE: Final = "RFQ Date"
_QUOTE_DUE: Final = "Quote Due"
_CURRENCY: Final = "Currency"
_NO: Final = "No."
_CODE: Final = "Customer Part No."
_DESCRIPTION: Final = "Description"
_QUANTITY: Final = "Quantity"
_UOM: Final = "UoM"
_TARGET_PRICE: Final = "Target Price"
_REQUIRED_DATE: Final = "Required Date"
_RFQ_COLUMNS: Final = frozenset(
    {_NO, _CODE, _DESCRIPTION, _QUANTITY, _UOM, _TARGET_PRICE, _REQUIRED_DATE}
)

_YCBG_NO: Final = "YCBG No."
_REPLY_DATE: Final = "Reply Date"
_BP_CODE: Final = "BP Code"
_SPEC_NO: Final = "Spec No."
_PRV_CODE: Final = "PRV Code"
_COPPER: Final = "Copper (kg/km)"
_REPLY_COLUMNS: Final = frozenset({_NO, _BP_CODE, _SPEC_NO, _PRV_CODE, _COPPER})

type _Position = tuple[int, int]
type _Cells = Mapping[_Position, object]


class ExcelRfqReader:
    """Implements `RfqReaderPort` for the mock customers' Excel layout."""

    def read(self, content: AttachmentContent) -> RfqDocument | None:
        sheet = _the_marked_sheet(content, _RFQ_NO)
        return None if sheet is None else _request(sheet)


class ExcelDesignReplyReader:
    """Implements `DesignReplyReaderPort` for the seller's Design reply form."""

    def read(self, content: AttachmentContent) -> DesignReplyDocument | None:
        sheet = _the_marked_sheet(content, _YCBG_NO)
        return None if sheet is None else _reply(sheet)


def _the_marked_sheet(content: AttachmentContent, marker: str) -> _Sheet | None:
    """The one sheet carrying ``marker``, or None when the file has none.

    Two such sheets are refused: which one is the document would be a guess.
    """
    attachment = content.attachment
    if not attachment.name.lower().endswith(_SUFFIX):
        return None
    sheets = _read_cells(content.data)
    if sheets is None:
        return None
    marked = []
    for title, cells in sheets:
        if not any(_is_text(value, marker) for value in cells.values()):
            continue
        sheet = _Sheet(attachment, title, cells, marker)
        try:
            sheet.anchor((1, 1))
        except ValidationError:
            # A sheet name Excel itself refuses (openpyxl only warns past
            # 31 characters) cannot be anchored, so nothing on it can be
            # shown to a reviewer: the file is not one this reader reads.
            return None
        marked.append(sheet)
    if not marked:
        return None
    if len(marked) > 1:
        second = marked[1]
        raise QuoteFileUnreadableError(
            f"more than one sheet carries {marker!r}", at=second.anchor(second.marker())
        )
    return marked[0]


def _read_cells(data: bytes) -> list[tuple[str, _Cells]] | None:
    """Every worksheet's non-empty cells, or None when the bytes are not a
    workbook this reader reads.

    Whatever zipfile, openpyxl or the XML parser raise on a malformed file, the
    answer is the same, so one handler takes all of it.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
        if len(entries) > _MAX_ENTRIES:
            return None
        if any(entry.filename == MACROS_PART for entry in entries):
            return None
        if sum(entry.file_size for entry in entries) > _MAX_UNCOMPRESSED_BYTES:
            return None
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets: list[tuple[str, _Cells]] = []
            for sheet in workbook.worksheets:
                cells: dict[_Position, object] = {}
                rows = sheet.iter_rows(max_row=_MAX_ROWS, max_col=_MAX_COLUMNS, values_only=True)
                for row, values in enumerate(rows, start=1):
                    for column, value in enumerate(values, start=1):
                        if value is not None and not (isinstance(value, str) and not value.strip()):
                            cells[(row, column)] = value
                sheets.append((sheet.title, cells))
            return sheets
        finally:
            workbook.close()
    except Exception:
        return None


def _is_text(value: object, text: str) -> bool:
    return isinstance(value, str) and value.strip() == text


def _text(raw: object) -> str:
    # A code may be a numeric cell; a float is not text anyone typed as a code.
    if isinstance(raw, bool) or not isinstance(raw, str | int):
        raise ValueError("not text")
    return str(raw).strip()


def _number(raw: object) -> Decimal:
    # openpyxl returns a numeric cell as int or float; str() gives back the
    # printed value. Text that looks like a number is refused, not parsed.
    if isinstance(raw, bool) or not isinstance(raw, int | float):
        raise ValueError("not a number")
    return Decimal(str(raw))


def _day(raw: object) -> date:
    # openpyxl returns a date cell as a datetime; a date written as text is refused.
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    raise ValueError("not a date")


def _line_no(raw: object) -> int:
    if isinstance(raw, bool):
        raise ValueError("not a line number")
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float) and raw.is_integer():
        return int(raw)
    raise ValueError("not a line number")


class _Sheet:
    """One worksheet's cells, read as the document ``marker`` labels."""

    def __init__(self, attachment: Attachment, title: str, cells: _Cells, marker: str) -> None:
        self._attachment = attachment
        self._title = title
        self._cells = cells
        self._marker = marker

    def anchor(self, at: _Position) -> SourceAnchor:
        row, column = at
        return SourceAnchor(
            attachment_id=self._attachment.attachment_id,
            attachment_sha256=self._attachment.sha256,
            cell_ref=f"{self._title}!{get_column_letter(column)}{row}",
        )

    def holds(self, at: _Position) -> bool:
        return at in self._cells

    def find(self, text: str, *, above: int | None = None) -> _Position | None:
        """Where ``text`` stands as a label, once, or None when it does not."""
        found = sorted(
            at
            for at, value in self._cells.items()
            if _is_text(value, text) and (above is None or at[0] < above)
        )
        if len(found) > 1:
            raise QuoteFileUnreadableError(
                f"the label {text!r} appears more than once", at=self.anchor(found[1])
            )
        return found[0] if found else None

    def label(self, text: str, *, above: int | None = None) -> _Position:
        found = self.find(text, above=above)
        if found is None:
            raise QuoteFileUnreadableError(
                f"the document has no {text!r}", at=self.anchor(self.marker())
            )
        return found

    def beside(self, text: str, *, above: int) -> _Position:
        row, column = self.label(text, above=above)
        return row, column + 1

    def marker(self) -> _Position:
        """The marker's label: what makes the sheet the document it is."""
        found = self.find(self._marker)
        if found is None:  # A sheet is built only once its marker has been seen.
            raise ValueError(f"sheet {self._title!r} carries no {self._marker!r}")
        return found

    def read[T](
        self, model: type[Sourced[T]], what: str, at: _Position, parse: Callable[[object], object]
    ) -> Sourced[T]:
        """The value at ``at``. Empty is refused: see `read_or_blank`."""
        if at not in self._cells:
            raise QuoteFileUnreadableError(f"{what} is empty", at=self.anchor(at))
        return self.read_or_blank(model, what, at, parse)

    def read_or_blank[T](
        self, model: type[Sourced[T]], what: str, at: _Position, parse: Callable[[object], object]
    ) -> Sourced[T]:
        """The value at ``at``, or None anchored at the empty cell, for a value
        the customer may leave blank and the case then asks for."""
        anchor = self.anchor(at)
        raw = self._cells.get(at)
        try:
            # ValidationError is a ValueError: a value of the wrong kind and one
            # the domain refuses (a zero quantity, an unknown unit) read alike.
            value = None if raw is None else parse(raw)
            return model.model_validate({"value": value, "source": anchor})
        except ValueError as exc:
            # The cell is named; its text is not repeated into the message,
            # because it is the sender's and the message reaches logs.
            raise QuoteFileUnreadableError(f"{what} cannot be read", at=anchor) from exc

    def optional[T](
        self, model: type[Sourced[T]], what: str, at: _Position, parse: Callable[[object], object]
    ) -> Sourced[T] | None:
        """The value at ``at``, or None when the cell is empty: a value the
        document does not have to state."""
        return self.read(model, what, at, parse) if at in self._cells else None

    def table(self, columns: frozenset[str]) -> tuple[int, dict[str, int]]:
        """The table's header row, and the column of each of its titles."""
        rows: dict[int, dict[str, int]] = {}
        for (row, column), value in sorted(self._cells.items()):
            if isinstance(value, str) and value.strip() in columns:
                rows.setdefault(row, {}).setdefault(value.strip(), column)
        for row in sorted(rows):
            if rows[row].keys() == columns:
                return row, rows[row]
        raise QuoteFileUnreadableError(
            "the document has no line table", at=self.anchor(self.marker())
        )

    def numbered_rows(self, header: int, number_column: int) -> list[tuple[int, int]]:
        """Each table row below ``header`` with its line number, down to the
        first row without one."""
        rows = []
        row = header + 1
        while (row, number_column) in self._cells:
            at = (row, number_column)
            if row == _MAX_ROWS:
                raise QuoteFileUnreadableError(
                    "the line table runs past the rows this reader reads", at=self.anchor(at)
                )
            try:
                rows.append((row, _line_no(self._cells[at])))
            except ValueError as exc:
                raise QuoteFileUnreadableError(
                    "the line number cannot be read", at=self.anchor(at)
                ) from exc
            row += 1
        if not rows:
            raise QuoteFileUnreadableError(
                "the document has no lines", at=self.anchor((header, number_column))
            )
        return rows

    def buyer(self) -> Sourced[Name] | None:
        """The buyer's name as the form prints it, or None when it prints none
        that reads as a name. Only ever compared with master data."""
        try:
            return self.optional(Sourced[Name], "the buyer", _BUYER_AT, _text)
        except QuoteFileUnreadableError:
            return None


def _request(sheet: _Sheet) -> RfqDocument:
    header, columns = sheet.table(_RFQ_COLUMNS)
    due = sheet.find(_QUOTE_DUE, above=header)
    items = []
    for row, line_no in sheet.numbered_rows(header, columns[_NO]):

        def at(title: str, row: int = row) -> _Position:
            return row, columns[title]

        try:
            items.append(
                RequestedItem(
                    line_no=line_no,
                    customer_item_code=sheet.optional(
                        Sourced[CustomerItemCode], "the customer part number", at(_CODE), _text
                    ),
                    description=sheet.read(
                        Sourced[Description], "the description", at(_DESCRIPTION), _text
                    ),
                    quantity=sheet.read_or_blank(
                        Sourced[Quantity | None], "the quantity", at(_QUANTITY), _number
                    ),
                    uom=sheet.read(Sourced[Uom], "the unit of measure", at(_UOM), _text),
                    target_price=sheet.optional(
                        Sourced[UnitPrice], "the target price", at(_TARGET_PRICE), _number
                    ),
                    needed_by=sheet.read_or_blank(
                        Sourced[date | None], "the required date", at(_REQUIRED_DATE), _day
                    ),
                )
            )
        except ValidationError as exc:
            raise QuoteFileUnreadableError(
                "the line cannot be read", at=sheet.anchor(at(_NO))
            ) from exc
    try:
        return RfqDocument(
            parser_version=PARSER_VERSION,
            buyer=sheet.buyer(),
            rfq_no=sheet.read(
                Sourced[DocumentNo], "the RFQ number", sheet.beside(_RFQ_NO, above=header), _text
            ),
            rfq_date=sheet.read(
                Sourced[date], "the RFQ date", sheet.beside(_RFQ_DATE, above=header), _day
            ),
            quote_due=(
                None
                if due is None
                else sheet.optional(Sourced[date], "the quote due date", (due[0], due[1] + 1), _day)
            ),
            currency=sheet.read(
                Sourced[Currency], "the currency", sheet.beside(_CURRENCY, above=header), _text
            ),
            items=tuple(items),
        )
    except ValidationError as exc:
        raise QuoteFileUnreadableError(
            "the request repeats a line number", at=sheet.anchor((header, columns[_NO]))
        ) from exc


def _reply(sheet: _Sheet) -> DesignReplyDocument:
    header, columns = sheet.table(_REPLY_COLUMNS)
    lines = []
    for row, line_no in sheet.numbered_rows(header, columns[_NO]):

        def at(title: str, row: int = row) -> _Position:
            return row, columns[title]

        try:
            lines.append(
                DesignReplyLine(
                    line_no=line_no,
                    bp_code=sheet.read(Sourced[BpCode], "the BP code", at(_BP_CODE), _text),
                    spec_no=sheet.read(Sourced[DocumentNo], "the spec number", at(_SPEC_NO), _text),
                    prv_code=sheet.optional(Sourced[PrvCode], "the PRV code", at(_PRV_CODE), _text),
                    copper_kg_per_km=sheet.optional(
                        Sourced[CopperWeight], "the copper weight", at(_COPPER), _number
                    ),
                )
            )
        except ValidationError as exc:
            raise QuoteFileUnreadableError(
                "the line cannot be read", at=sheet.anchor(at(_NO))
            ) from exc
    try:
        return DesignReplyDocument(
            ycbg_no=sheet.read(
                Sourced[DocumentNo], "the YCBG number", sheet.beside(_YCBG_NO, above=header), _text
            ),
            reply_date=sheet.read(
                Sourced[date], "the reply date", sheet.beside(_REPLY_DATE, above=header), _day
            ),
            lines=tuple(lines),
        )
    except ValidationError as exc:
        raise QuoteFileUnreadableError(
            "the reply answers a line twice", at=sheet.anchor((header, columns[_NO]))
        ) from exc
