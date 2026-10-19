"""`CatalogTemplatePort`: a list's template as an ``.xlsx`` (ticket 41).

Two sheets: "Dữ liệu", the header row in the list's column order, which the
reader reads first; and "Hướng dẫn", each column's meaning, whether the file
must have it, and the words it takes. Written, never read, in this process:
nothing of a customer's file is opened here.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from dw_sales.domain.master_data import ListSpec

_DATA = "Dữ liệu"
_GUIDE = "Hướng dẫn"


class XlsxCatalogTemplates:
    """Implements `CatalogTemplatePort`."""

    def workbook(self, spec: ListSpec) -> bytes:
        book = Workbook()
        data = book.active
        assert data is not None
        data.title = _DATA
        data.append([column.header for column in spec.columns])
        for cell in data[1]:
            cell.font = Font(bold=True)
        for index, column in enumerate(spec.columns, start=1):
            data.column_dimensions[get_column_letter(index)].width = max(14, len(column.header) + 4)
        data.freeze_panes = "A2"
        guide = book.create_sheet(_GUIDE)
        guide.append(["Cột", "Bắt buộc có cột", "Cách ghi"])
        for cell in guide[1]:
            cell.font = Font(bold=True)
        for column in spec.columns:
            guide.append([column.header, "Có" if column.required else "Không", column.hint])
        guide.append([])
        guide.append(
            [
                "Lưu ý",
                "",
                f"Mỗi dòng ở trang '{_DATA}' là một bản ghi {spec.label.lower()}. Tệp nhập"
                " thay cả danh mục khi được duyệt: dòng không có trong tệp sẽ bị bỏ.",
            ]
        )
        guide.column_dimensions["A"].width = 32
        guide.column_dimensions["B"].width = 16
        guide.column_dimensions["C"].width = 90
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()
