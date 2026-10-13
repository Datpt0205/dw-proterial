"""`AttachmentTextPort`: the text of an attachment a model may read (ticket 29).

A PDF's text layer, page by page, and a workbook's visible cells, sheet by
sheet, each as one `TextPart` under the caps the order rules set (spec
decision 13). Parsing only, and no value is read here: the model reads, code
verifies (`domain.rfq_reading`).

**Hidden content is not sent.** A hidden sheet, row or column, a cell whose
text is drawn in its own fill colour, and a formula with no cached value are
left out: what a person looking at the file would not see is not part of the
request, and it is the injection channel (`ai-automation.md`, "What stays
human"). A scan (no text layer), an encrypted or macro-enabled file, or one
over the caps gives no text at all, and the request goes to a person.

A sheet's text is its rows, top to bottom, each its visible cells left to
right joined by `` | ``; every cell's place in that text is kept, so a value
the model points at is anchored at its cell.
"""

from __future__ import annotations

import asyncio
import io
import re
from collections.abc import Sequence
from datetime import date, datetime
from typing import Final

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from dw_sales.adapters.readers.documents import sniff
from dw_sales.adapters.readers.excel import font_matches_fill, open_for_viewing
from dw_sales.domain.messages import AttachmentContent
from dw_sales.domain.order_checks import IntakeCaps
from dw_sales.domain.rfq_reading import TextPart, normalised

# What is read of each sheet, whatever dimensions it declares (as the
# fixed-label reader).
_MAX_ROWS: Final = 1000
_MAX_COLUMNS: Final = 26
_CELL_JOIN: Final = " | "
_SPACES = re.compile(r"[ \t]+")


class AttachmentTexts:
    """Implements `AttachmentTextPort` for text PDFs and .xlsx workbooks."""

    async def text_parts(self, content: AttachmentContent, caps: IntakeCaps) -> Sequence[TextPart]:
        kind = sniff(content.data)
        if content.attachment.size > caps.max_attachment_bytes:
            return ()
        # Parsing is CPU work on a file of the sender's choosing: off the loop.
        if kind == "pdf":
            return await asyncio.to_thread(_pdf_parts, content, caps)
        if kind == "xlsx":
            return await asyncio.to_thread(_sheet_parts, content, caps)
        return ()


def _pdf_parts(content: AttachmentContent, caps: IntakeCaps) -> list[TextPart]:
    attachment = content.attachment
    try:
        pdf = PdfReader(io.BytesIO(content.data))
        if pdf.is_encrypted or len(pdf.pages) > caps.max_pages:
            return []
        parts = []
        for number, page in enumerate(pdf.pages, start=1):
            text = _clean(page.extract_text())
            if text:
                parts.append(
                    TextPart(
                        part_id=f"{attachment.attachment_id}#p{number}",
                        text=text,
                        attachment_id=attachment.attachment_id,
                        sha256=attachment.sha256,
                        page=number,
                    )
                )
        return parts
    except Exception:  # any parser failure on a sender's file: no text
        return []


def _sheet_parts(content: AttachmentContent, caps: IntakeCaps) -> list[TextPart]:
    attachment = content.attachment
    try:
        values, _ = open_for_viewing(content.data, caps)
    except ValueError:
        return []
    parts = []
    for sheet in values.worksheets:
        if sheet.sheet_state != "visible":
            continue
        hidden_columns = {
            index
            for index in range(1, _MAX_COLUMNS + 1)
            if sheet.column_dimensions[get_column_letter(index)].hidden
        }
        lines: list[str] = []
        cells: list[tuple[int, int, str]] = []
        offset = 0
        rows = min(sheet.max_row, _MAX_ROWS)
        columns = min(sheet.max_column, _MAX_COLUMNS)
        for row in range(1, rows + 1):
            if sheet.row_dimensions[row].hidden:
                continue
            pieces: list[str] = []
            at = offset
            for column in range(1, columns + 1):
                if column in hidden_columns:
                    continue
                cell = sheet.cell(row=row, column=column)
                text = _cell_text(cell)
                if not text:
                    continue
                if pieces:
                    at += len(_CELL_JOIN)
                ref = f"{sheet.title}!{get_column_letter(column)}{row}"
                cells.append((at, at + len(text), ref))
                pieces.append(text)
                at += len(text)
            if pieces:
                line = _CELL_JOIN.join(pieces)
                lines.append(line)
                offset += len(line) + 1
        if lines:
            parts.append(
                TextPart(
                    part_id=f"{attachment.attachment_id}!{sheet.title}",
                    text="\n".join(lines),
                    attachment_id=attachment.attachment_id,
                    sha256=attachment.sha256,
                    cells=tuple(cells),
                )
            )
    return parts


def _cell_text(cell: Cell | MergedCell) -> str:
    """What a person sees in the cell, NFKC, on one line; empty for an empty
    cell (a formula never computed reads as one), a merged cell's covered
    part, and a cell whose text is drawn in its fill colour."""
    value = cell.value
    if value is None or not isinstance(cell, Cell) or font_matches_fill(cell):
        return ""
    if isinstance(value, datetime):
        midnight = value.time() == datetime.min.time()
        text = value.date().isoformat() if midnight else value.isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float):
        text = format(value, ".12g")
    else:
        text = str(value)
    return _SPACES.sub(" ", normalised(" ".join(text.split()))).strip().replace("|", "/")


def _clean(text: str) -> str:
    lines = [_SPACES.sub(" ", normalised(line)).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)
