"""Ticket 15's header path on a fixture: a PO whose header value is flagged.

No PO of the mailbox flags a header value, so the correction of one was
covered by hand-built cases only. The fixture here is M01's workbook under
its own PO number with the PO date typed white on white (`font_matches_fill`),
sent as an eval-only message (`HEADER_MESSAGE`, beside the mailbox, never in
it: the demo, the browser walk and the docs count its 45 messages). It is
read by the real readers and decided through the real services:

- it raises the PO-level `value_uncertain` on that value alone, blocking;
- `prepare` is refused naming it, even with every source opened;
- a field the finding did not flag is refused;
- Sales types the date; the finding is corrected, the header is entered
  with the typed date and the date as read is kept beside it;
- and the PO, otherwise clean, is then prepared.

Details name the message, the finding and the field, never a value read.
"""

from __future__ import annotations

import asyncio
import io
import tempfile
from collections.abc import Awaitable
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from dw_evals.graders import GraderContext, GradeResult
from dw_kernel.errors import ConflictError, DWError
from dw_sales.adapters.mock.fixtures import ATTACHMENTS_DIR
from dw_sales.domain.orders import FindingCode, OrderCase, RegionFlag
from dw_sales.evals.graders import AN, CALLERS
from dw_sales.evals.world import SampleWith, World, open_world

HEADER_MESSAGE = "M92"
WHITE = "FFFFFFFF"


def header_flagged_po(po_no: str, cell: str) -> bytes:
    """M01's workbook under ``po_no``, its ``cell`` typed in its own fill
    colour: there, and invisible to a person reading the sheet."""
    workbook = load_workbook(ATTACHMENTS_DIR / "M01_VLX-PO-2609-0118.xlsx")
    sheet = workbook["PO"]
    sheet["B4"] = po_no
    sheet[cell].font = Font(color=WHITE)
    sheet[cell].fill = PatternFill(fill_type="solid", fgColor=WHITE)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _message(po_no: str, path: Path) -> dict[str, Any]:
    return {
        "message_id": HEADER_MESSAGE,
        "sender": "an.trinh@velatrix.example",
        "subject": f"PO {po_no}",
        "received_at": "2026-09-21T10:30:00+07:00",
        "body": f"Chào chị Vy,\n\nVelatrix gửi PO {po_no} đính kèm.\n\nTrân trọng,\nTrịnh Mỹ An",
        "authentication": {"spf": "pass", "dkim": "pass", "dmarc": "pass"},
        "attachments": [
            {
                "attachment_id": f"{HEADER_MESSAGE}-A1",
                "name": f"{po_no}.xlsx",
                "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "file": str(path),
            }
        ],
    }


def grade_header_corrected(
    ctx: GraderContext, input_data: dict[str, Any], expected: dict[str, Any]
) -> GradeResult:
    """A header value a person would not see is decided value by value
    (ticket 15): blocking until Sales types it, then entered as typed."""
    po_no: str = input_data["po_no"]
    with tempfile.TemporaryDirectory() as scratch:
        path = Path(scratch) / f"{po_no}.xlsx"
        path.write_bytes(header_flagged_po(po_no, input_data["cell"]))
        # The world reads the attachment again to serve its sources: it
        # lives as long as the sample set does.
        with SampleWith([_message(po_no, path)]) as sample:
            world = open_world(ctx.repo_root, sample, CALLERS)
            return asyncio.run(_header_corrected(world, input_data, expected))


async def _refused(call: Awaitable[object]) -> bool:
    try:
        await call
    except DWError:  # 403, 404, 409 and 422 alike: the call was refused
        return True
    return False


async def _header_corrected(
    world: World, input_data: dict[str, Any], expected: dict[str, Any]
) -> GradeResult:
    field: str = input_data["field"]
    typed: str = input_data["typed"]
    key: str = expected["finding"]
    ((disposition, case),) = (await world.process_each(AN, [HEADER_MESSAGE])).values()
    if not isinstance(case, OrderCase):
        return GradeResult.fail("the PO opened no order case", disposition=disposition.kind)
    finding = next((f for f in case.findings if f.key == key), None)
    if finding is None or not finding.blocking:
        keys = [f.key for f in case.findings]
        return GradeResult.fail(f"{key} is not a blocking finding", keys=keys)
    problems: list[str] = []
    flagged = sorted(case.flagged_fields(finding))
    flags = [f.flag for f in case.header.flags]
    if flagged != [field] or flags != [RegionFlag(expected["flag"])]:
        problems.append(f"{key} flags {flagged}, not {field} alone")
    uncertain = [f.key for f in case.findings if f.code is FindingCode.VALUE_UNCERTAIN]
    if uncertain != [key]:
        problems.append(f"value_uncertain raised on {uncertain}")
    commands, context = world.services.order_commands, AN.context()
    await world.open_sources(AN, case.case_id)
    try:
        await commands.prepare(context, case.case_id, case_version=case.case_version)
    except ConflictError as refused:
        named = {str(k) for k in cast(list[object], refused.details.get("open_findings", []))}
        if key not in named:
            problems.append(f"prepare refused without naming {key}")
    else:
        problems.append(f"prepared with {key} open")

    def correct(name: str, value: str) -> Awaitable[object]:
        return commands.correct_value(
            context,
            case.case_id,
            key,
            case_version=world.order(case.case_id).case_version,
            field=name,
            value=value,
        )

    # A value the other field could hold, so only its not being flagged refuses it.
    other, value = input_data["unflagged"]
    if not await _refused(correct(other, value)):
        problems.append(f"{other}, which {key} did not flag, took a value")
    if await _refused(correct(field, typed)):
        return GradeResult.fail(f"{field} could not be typed", problems=problems)
    corrected = world.order(case.case_id)
    if next(f for f in corrected.findings if f.key == key).is_open:
        problems.append(f"{key} stayed open once {field} was typed")
    entered = str(getattr(corrected.effective_header, field))
    read = str(getattr(corrected.header, field))
    if entered != typed or read == typed:
        problems.append(f"{field} is not entered as typed beside the value as read")
    await world.open_sources(AN, case.case_id)
    if await _refused(commands.prepare(context, case.case_id, case_version=corrected.case_version)):
        problems.append("the corrected PO could not be prepared")
    if problems:
        return GradeResult.fail(
            "a flagged header value was not decided by field", problems=problems
        )
    return GradeResult.ok(message=HEADER_MESSAGE, finding=key, field=field)


GRADERS = {"sales.header_corrected": grade_header_corrected}
