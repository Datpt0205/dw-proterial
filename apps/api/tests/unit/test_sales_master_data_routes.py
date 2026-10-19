"""The master-data import routes (Sales ticket 41), over the in-memory Sales
services: what the API answers, and what it refuses, before any database."""

from __future__ import annotations

import base64
import io
import uuid
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

from dw_api.exception_handlers import register_exception_handlers
from dw_platform.application.access_context import AccessContext
from dw_sales.domain.master_data import SPECS, CatalogList
from dw_sales.evals.graders import AN, CALLERS, GIANG, OUTSIDER, READER, Caller
from dw_sales.evals.world import BETA, SampleSet, World, open_world
from dw_sales.presentation.routes import build_router

pytestmark = pytest.mark.unit

REPO = Path(__file__).resolve().parents[4]
BASE = "/api/v1/sales/master-data/imports"


class _Recorder:
    """The idempotency operation, as the app's answers a first request."""

    async def record[M: BaseModel](self, result: M, *, status_code: int = 200) -> M:
        return result


def _app(world: World, caller: Caller, scope: Any = None) -> FastAPI:
    def context() -> AccessContext:
        return caller.context(scope) if scope is not None else caller.context()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(
        build_router(world.services, access_context=context, idempotency=lambda: _Recorder())
    )
    return app


def _client(app: FastAPI) -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")


def _lme_file() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Dữ liệu"
    sheet.append([c.header for c in SPECS[CatalogList.LME].columns])
    sheet.append(["09/2026", 9600])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _body(data: bytes, name: str = "lme.xlsx") -> dict[str, str]:
    return {"file_name": name, "content_b64": base64.b64encode(data).decode()}


@pytest.fixture
def world() -> World:
    return open_world(REPO, SampleSet.of(REPO, None), (*CALLERS, OUTSIDER))


async def test_an_upload_is_queued_previewed_and_submitted(world: World) -> None:
    async with _client(_app(world, GIANG)) as client:
        queued = await client.post(f"{BASE}/lme", json=_body(_lme_file()))
        assert queued.status_code == 202, queued.text
        snapshot_id = queued.json()["snapshot_id"]
        assert queued.json()["status"] == "queued"
        await world.read_imports()
        preview = await client.get(f"{BASE}/snapshots/{snapshot_id}")
        assert preview.status_code == 200
        assert preview.json()["snapshot"]["status"] == "parsed"
        assert preview.json()["compared_with"]["source"] == "mock"
        submitted = await client.post(f"{BASE}/snapshots/{snapshot_id}/submit")
        assert submitted.status_code == 200
        assert submitted.json()["approval_id"] is not None
        state = await client.get(BASE)
        lme = next(s for s in state.json()["lists"] if s["catalog_list"] == "lme")
        assert lme["open"]["status"] == "submitted"


async def test_the_template_downloads_as_a_workbook_with_vietnamese_headers(world: World) -> None:
    async with _client(_app(world, AN)) as client:
        response = await client.get(f"{BASE}/customers/template")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="mau-customers.xlsx"'
    sheet = load_workbook(io.BytesIO(response.content))["Dữ liệu"]
    assert [c.value for c in sheet[1]][:2] == ["Mã khách hàng", "Tên khách hàng"]


async def test_a_reader_is_refused_and_another_tenant_finds_nothing(world: World) -> None:
    async with _client(_app(world, GIANG)) as client:
        queued = await client.post(f"{BASE}/lme", json=_body(_lme_file()))
    snapshot_id = queued.json()["snapshot_id"]
    await world.read_imports()
    async with _client(_app(world, READER)) as client:
        refused = await client.post(f"{BASE}/lme", json=_body(_lme_file()))
        assert refused.status_code == 403
        assert (await client.get(f"{BASE}/customers/template")).status_code == 403
    async with _client(_app(world, OUTSIDER, BETA)) as client:
        assert (await client.get(f"{BASE}/snapshots/{snapshot_id}")).status_code == 404
        assert (await client.post(f"{BASE}/snapshots/{snapshot_id}/submit")).status_code == 404
        assert (await client.get(f"{BASE}/snapshots/{uuid.uuid4()}")).status_code == 404


async def test_a_bad_upload_is_refused_with_a_reason(world: World) -> None:
    async with _client(_app(world, GIANG)) as client:
        wrong_type = await client.post(f"{BASE}/lme", json=_body(b"%PDF-1", "a.pdf"))
        assert wrong_type.status_code == 422
        assert wrong_type.json()["message"] == "chỉ nhận tệp .xlsx, .xls hoặc .csv"
        garbled = await client.post(
            f"{BASE}/lme", json={"file_name": "a.xlsx", "content_b64": "!!not base64!!"}
        )
        assert garbled.status_code == 422
        unknown = await client.post(f"{BASE}/bravo_orders", json=_body(_lme_file()))
        assert unknown.status_code == 422
