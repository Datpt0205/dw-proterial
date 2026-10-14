"""Artifacts through the API (ticket 06): rendered from the case by a PIC,
listed with their template, case version and hash, and downloaded only in
their state, for the version they were rendered at, by a caller who may see
what they carry. Before its state 409, after an edit 409, as the viewer 403,
another case's or another tenant's 404.
"""

from __future__ import annotations

import email
import email.message
import email.policy
import hashlib
import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook
from sales_api_harness import AN, BAO, DIEU, GIANG, HA, Api, Persona, user_id
from sales_flow import (
    approve,
    cross_checked,
    decide,
    order_of,
    prepared,
    process,
    quote,
    replied,
    submitted,
    uploaded,
)

from dw_kernel.ids import TenantId, WorkspaceId
from dw_platform.adapters.persistence.repositories import (
    SqlApprovalRepository,
    SqlAuditRepository,
)
from dw_platform.testing.seed_env import sid
from dw_sales.adapters.persistence.uow import SqlSalesUnitOfWorkFactory
from dw_sales.application.artifacts_service import ORDER_GATES, QUOTE_GATES
from dw_sales.application.case_store import ArtifactRecord
from dw_sales.application.ports import SalesScope
from dw_sales.domain.dispositions import CaseKind

pytestmark = pytest.mark.integration

ALPHA = SalesScope(
    TenantId(sid("tenant", "tenant-alpha")), WorkspaceId(sid("workspace", "tenant-alpha:main"))
)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _artifact(api: Api, kind: str, case_kind: CaseKind, case: dict[str, Any]) -> str:
    """A record written straight to the store, as no render could before its state."""
    record = ArtifactRecord(
        artifact_id=uuid.uuid4(),
        case_kind=case_kind,
        case_id=uuid.UUID(case["case_id"]),
        case_version=case["case_version"],
        kind=kind,
        template_ref="sales_test@1.0.0",
        sha256=hashlib.sha256(b"file").hexdigest(),
        content_type="application/octet-stream",
        size_bytes=4,
        created_by=user_id(AN),
        created_at=datetime.now(UTC),
    )
    uow = SqlSalesUnitOfWorkFactory(
        api.sessions, audit=SqlAuditRepository, approvals=SqlApprovalRepository
    )
    async with uow(ALPHA) as work:
        await work.artifacts.add(record)
        await work.commit()
    api.artifacts.objects[record.object_key(ALPHA)] = b"file"
    return str(record.artifact_id)


async def _render(
    persona: Persona, path: str, case: dict[str, Any], kind: str, *, key: str | None = None
) -> list[dict[str, Any]]:
    response = await persona.post(
        f"/{path}/{case['case_id']}/artifacts",
        {"kind": kind, "case_version": case["case_version"]},
        key=key,
    )
    assert response.status_code == 200, response.text
    artifacts: list[dict[str, Any]] = response.json()["artifacts"]
    return artifacts


# ------------------------------------------------------- before its state --


@pytest.mark.parametrize("kind", sorted(ORDER_GATES))
async def test_an_order_artifact_before_its_state_is_refused(api: Api, kind: str) -> None:
    an = api.as_(AN)
    case = await order_of(an, "M01")  # in review: no artifact's state yet
    artifact = await _artifact(api, kind, CaseKind.ORDER, case)

    response = await an.get(f"/orders/{case['case_id']}/artifacts/{artifact}")
    rendering = await an.post(
        f"/orders/{case['case_id']}/artifacts",
        {"kind": kind, "case_version": case["case_version"]},
    )

    assert response.status_code == 409
    assert response.json()["details"]["kind"] == kind
    assert rendering.status_code == 409, rendering.text


@pytest.mark.parametrize("kind", sorted(QUOTE_GATES))
async def test_a_quote_artifact_before_its_state_is_refused(api: Api, kind: str) -> None:
    dieu = api.as_(DIEU)
    opened = await process(dieu, "M10")  # received: no artifact's state yet
    case = await quote(dieu, opened["case_id"])
    artifact = await _artifact(api, kind, CaseKind.QUOTE, case)

    response = await dieu.get(f"/quotes/{case['case_id']}/artifacts/{artifact}")
    rendering = await dieu.post(
        f"/quotes/{case['case_id']}/artifacts",
        {"kind": kind, "case_version": case["case_version"]},
    )

    assert response.status_code == 409
    assert rendering.status_code == 409, rendering.text


# ------------------------------------------------------- the upload file --


async def test_an_renders_the_upload_file_and_downloads_exactly_what_was_recorded(
    api: Api,
) -> None:
    an = api.as_(AN)
    ready = await prepared(an, await order_of(an, "M01"))

    (rendered,) = await _render(an, "orders", ready, "bravo_upload")
    listed = await an.get(f"/orders/{ready['case_id']}/artifacts")
    served = await an.get(f"/orders/{ready['case_id']}/artifacts/{rendered['artifact_id']}")

    assert listed.status_code == 200
    body = listed.json()
    assert body["case_version"] == ready["case_version"]
    assert "bravo_upload" in body["available"]
    (record,) = body["artifacts"]
    assert record == rendered
    assert (record["template_ref"], record["case_version"]) == (
        "sales_bravo_upload@1.1.0",
        ready["case_version"],
    )
    assert served.status_code == 200
    assert hashlib.sha256(served.content).hexdigest() == record["sha256"]
    assert served.headers["content-type"] == XLSX
    assert served.headers["cache-control"] == "no-store"
    assert (
        served.headers["content-disposition"]
        == f'attachment; filename="bravo_upload-{ready["case_id"]}.xlsx"'
    )
    sheet = load_workbook(io.BytesIO(served.content)).worksheets[0]
    assert [line["line_no"] for line in ready["lines"]] == [
        row[4] for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


async def test_a_line_split_over_several_dates_keeps_each_through_the_store_and_the_file(
    api: Api,
) -> None:
    """Ticket 21, read back from PostgreSQL: M34 prints lines 1 and 3 over
    two and three dates, and the order and its upload file keep every one."""
    an = api.as_(AN)
    case = await order_of(an, "M34")

    assert [(line["printed_schedules"], len(line["schedules"])) for line in case["lines"]] == [
        (True, 2),
        (False, 1),
        (True, 3),
    ]
    ready = await prepared(an, case)
    (rendered,) = await _render(an, "orders", ready, "bravo_upload")
    served = await an.get(f"/orders/{ready['case_id']}/artifacts/{rendered['artifact_id']}")

    assert served.status_code == 200
    sheet = load_workbook(io.BytesIO(served.content)).worksheets[0]
    rows = [
        (r[4], Decimal(str(r[8])), r[13].date())  # Dong, SoLuong, NgayYeuCau
        for r in sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert rows == [
        (line["line_no"], Decimal(s["quantity"]), date.fromisoformat(s["requested_date"]))
        for line in ready["lines"]
        for s in line["schedules"]
    ]
    assert len(rows) == 6


async def test_a_split_line_is_confirmed_per_schedule_through_the_store_and_the_draft(
    api: Api,
) -> None:
    """Ticket 35, read back from PostgreSQL: one date for a line M34 splits is
    refused; a date per schedule is stored, shown per schedule and drafted
    to the customer schedule by schedule."""
    an, dieu = api.as_(AN), api.as_(DIEU)
    crossed = await cross_checked(dieu, await uploaded(an, await order_of(an, "M34")))
    path = f"/orders/{crossed['case_id']}/confirm"

    one_date = await an.post(
        path,
        {
            "case_version": crossed["case_version"],
            "delivery_dates": [
                {"line_no": line["line_no"], "confirmed_date": "2026-11-27"}
                for line in crossed["lines"]
            ],
        },
    )
    per_schedule = {
        1: ["2026-11-18", "2026-12-02"],
        2: ["2026-11-27"],
        3: ["2026-11-27", "2026-12-11", "2026-12-29"],
    }
    confirmed = await decide(
        an,
        crossed,
        "confirm",
        {
            "delivery_dates": [
                {"line_no": 2, "confirmed_date": "2026-11-27"},
                {"line_no": 1, "schedule_dates": per_schedule[1]},
                {"line_no": 3, "schedule_dates": per_schedule[3]},
            ]
        },
    )

    assert one_date.status_code == 422, one_date.text
    assert one_date.json()["details"]["lines"] == [1, 3]
    assert confirmed["status"] == "confirmed"
    by_line = {line["line_no"]: line for line in confirmed["lines"]}
    assert [s["confirmed_delivery_date"] for s in by_line[1]["schedules"]] == per_schedule[1]
    assert [s["confirmed_delivery_date"] for s in by_line[3]["schedules"]] == per_schedule[3]
    assert by_line[1]["confirmed_delivery_date"] is None
    assert by_line[2]["confirmed_delivery_date"] == "2026-11-27"
    (draft,) = await _render(an, "orders", confirmed, "confirmation_draft")
    served = await an.get(f"/orders/{confirmed['case_id']}/artifacts/{draft['artifact_id']}")
    assert served.status_code == 200
    message = email.message_from_bytes(served.content, policy=email.policy.default)
    assert isinstance(message, email.message.EmailMessage)
    part = message.get_body(("plain",))
    assert part is not None
    body = str(part.get_content())
    assert "theo 2 đợt: 6100 m ngày 18/11/2026; 6100 m ngày 02/12/2026" in body
    assert "theo 3 đợt: 1500 m ngày 27/11/2026; 900 m ngày 11/12/2026; 600 m ngày 29/12/2026" in (
        body
    )


async def test_the_upload_file_is_refused_once_an_edit_bumped_the_case(api: Api) -> None:
    """Recording the Bravo entry keeps the upload file's gate open and moves
    the case on: the file rendered before it is stale, the one rendered after
    is served."""
    an = api.as_(AN)
    ready = await prepared(an, await order_of(an, "M01"))
    (before,) = await _render(an, "orders", ready, "bravo_upload")
    entered = await decide(an, ready, "bravo-entry", {"so_no": "SO26-1001", "entry_compared": True})

    stale = await an.get(f"/orders/{ready['case_id']}/artifacts/{before['artifact_id']}")
    (after,) = await _render(an, "orders", entered, "bravo_upload")
    fresh = await an.get(f"/orders/{ready['case_id']}/artifacts/{after['artifact_id']}")

    assert stale.status_code == 409
    assert stale.json()["details"]["reason"] == "stale_artifact"
    assert fresh.status_code == 200
    assert after["case_version"] == entered["case_version"] > before["case_version"]


async def test_a_render_with_a_stale_version_is_refused(api: Api) -> None:
    an = api.as_(AN)
    case = await order_of(an, "M01")
    ready = await prepared(an, case)

    response = await an.post(
        f"/orders/{ready['case_id']}/artifacts",
        {"kind": "bravo_upload", "case_version": case["case_version"]},
    )

    assert response.status_code == 409


async def test_a_retried_render_answers_the_same_files(api: Api) -> None:
    an = api.as_(AN)
    ready = await prepared(an, await order_of(an, "M01"))

    first = await _render(an, "orders", ready, "bravo_upload", key="render-once")
    again = await _render(an, "orders", ready, "bravo_upload", key="render-once")

    assert first == again
    assert len(api.artifacts.objects) == 1


async def test_the_viewer_and_another_tenant_get_nothing(api: Api) -> None:
    an = api.as_(AN)
    ready = await prepared(an, await order_of(an, "M01"))
    (rendered,) = await _render(an, "orders", ready, "bravo_upload")
    path = f"/orders/{ready['case_id']}/artifacts"

    for persona, status in ((api.as_(HA), 403), (api.as_(BAO, tenant="tenant-beta"), 404)):
        assert (await persona.get(path)).status_code == status
        assert (await persona.get(f"{path}/{rendered['artifact_id']}")).status_code == status
        rendering = await persona.post(
            path, {"kind": "bravo_upload", "case_version": ready["case_version"]}
        )
        assert rendering.status_code == status


async def test_another_cases_artifact_is_not_found(api: Api) -> None:
    an = api.as_(AN)
    first = await prepared(an, await order_of(an, "M01"))
    second = await order_of(an, "M06")
    (rendered,) = await _render(an, "orders", first, "bravo_upload")

    response = await an.get(f"/orders/{second['case_id']}/artifacts/{rendered['artifact_id']}")

    assert response.status_code == 404


async def test_the_cross_check_sheet_waits_for_the_bravo_entry(api: Api) -> None:
    an = api.as_(AN)
    case = await uploaded(an, await order_of(an, "M01"))

    (sheet,) = await _render(an, "orders", case, "cross_check_sheet")

    assert sheet["template_ref"] == "sales_documents.cross_check_sheet.vi@1.2.0"


# ----------------------------------------------------------- the quotation --


async def test_giang_downloads_the_approved_quotation_and_the_send_draft_attaches_it(
    api: Api,
) -> None:
    """Spec "Done when": Giang approves M10 and downloads the approved
    quotation document; its hash is the one the approval names."""
    dieu, giang = api.as_(DIEU), api.as_(GIANG)
    pending = await submitted(dieu, await replied(dieu, "M10", "YCBG-2609-030", "M25"))
    previews = await _render(giang, "quotes", pending, "quotation_preview")
    assert {p["content_type"] for p in previews} == {XLSX, "application/pdf"}
    approved = await approve(giang, pending)
    assert approved.status_code == 200, approved.text
    final = await quote(giang, pending["case_id"])

    files = await _render(giang, "quotes", final, "quotation_final")
    (draft,) = await _render(dieu, "quotes", final, "send_draft")
    served = {
        f["content_type"]: await giang.get(
            f"/quotes/{final['case_id']}/artifacts/{f['artifact_id']}"
        )
        for f in files
    }

    assert {f["template_ref"] for f in files} == {"sales_documents.quotation.ja@1.2.0"}
    assert all(response.status_code == 200 for response in served.values())
    assert final["approval"]["document_sha256"].encode() in b"".join(
        str(c.value).encode()
        for row in load_workbook(io.BytesIO(served[XLSX].content)).worksheets[0].iter_rows()
        for c in row
    )
    assert draft["template_ref"] == "sales_emails.send_draft.ja@1.1.0"
    # A preview is for the pricer and approvers, and only while it waits.
    stale = await giang.get(f"/quotes/{final['case_id']}/artifacts/{previews[0]['artifact_id']}")
    assert stale.status_code == 409


async def test_a_pic_who_is_neither_pricer_nor_approver_gets_no_preview(api: Api) -> None:
    dieu, an = api.as_(DIEU), api.as_(AN)
    pending = await submitted(dieu, await replied(dieu, "M10", "YCBG-2609-030", "M25"))

    response = await an.post(
        f"/quotes/{pending['case_id']}/artifacts",
        {"kind": "quotation_preview", "case_version": pending["case_version"]},
    )

    assert response.status_code == 409
